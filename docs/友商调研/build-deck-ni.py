# -*- coding: utf-8 -*-
"""NI 如何设计 ATE：两页图文并茂（含 NI 官网采集图片）"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "ni_img")
OUT = os.path.join(HERE, "NI-ATE软件栈.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x18, 0x18, 0x1A)
MUT = RGBColor(0x6E, 0x6C, 0x68)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x0F, 0x4C, 0x81)
GREEN = RGBColor(0x2F, 0x6B, 0x57)
AMBER = RGBColor(0xB0, 0x76, 0x1C)
PURPLE = RGBColor(0x5C, 0x3A, 0x7A)
TEAL = RGBColor(0x20, 0x60, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD6, 0xD4, 0xCE)
SOFT = RGBColor(0xF9, 0xF9, 0xF7)
CREAM = RGBColor(0xFF, 0xF8, 0xEC)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _ea(run, size_pt, bold, color):
    f = run.font
    f.size = Pt(size_pt)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(sl, x, y, w, h, text, size=11, bold=False, color=INK,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        _ea(r, size, bold, color)
    return tb


def box(sl, x, y, w, h, fill, line=None, radius=0.10, lw=1.1):
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                              Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.text_frame.text = ""
    return shp


def header(sl, title, sub, page=None):
    box(sl, 0, 0, SW, 0.10, RED, None, 0.0)
    txt(sl, 0.55, 0.30, 11.9, 0.50, title, 23, True, INK)
    txt(sl, 0.57, 0.88, 11.6, 0.30, sub, 11, False, MUT)
    if page:
        txt(sl, 12.42, 7.14, 0.45, 0.22, f"{page}", 8.5, False, MUT, PP_ALIGN.RIGHT)


def footer(sl, s):
    txt(sl, 0.55, 7.12, 11.3, 0.24, s, 7.5, False, RGBColor(0x98, 0x96, 0x90))


def pic(sl, name, x, y, w, cap=None, cap_h=0.26):
    p = os.path.join(IMG, name)
    im = Image.open(p)
    h = w * im.size[1] / im.size[0]
    sl.shapes.add_picture(p, Inches(x), Inches(y), width=Inches(w))
    if cap:
        txt(sl, x, y + h + 0.05, w, cap_h, cap, 7, False, RGBColor(0x8A, 0x88, 0x82))
    return h


# ─────────────────────────── P1：软件栈 + 官网图片
s1 = prs.slides.add_slide(BLANK)
header(s1, "NI 怎么设计 ATE：五层软件，各管一件事",
       "能力来自「叠层」而不是某个万能工具——从机箱配置到数据归档，每层都有专门的人负责的软件", 1)

layers = [
    ("L5 数据与运维", "机群与测试数据管理、离线分析与报告",
     ["SystemLink", "DIAdem"], PURPLE),
    ("L4 执行与调度", "序列引擎、流程模型、site/pin map 与 binning",
     ["TestStand", "TSM 半导体模块"], RED),
    ("L3 测量与交互", "测量代码本体；仪器软面板与交互调试",
     ["LabVIEW", "InstrumentStudio"], BLUE),
    ("L2 路由与驱动", "开关矩阵路径；会话层与仪器/传感器驱动",
     ["Switch Executive", "NI-VISA / 488.2", "NI-DAQmx"], GREEN),
    ("L1 平台与硬件", "机箱、模块、时钟与触发的发现、配置与自检",
     ["PXI 机箱与模块", "PXI Platform Services", "NI MAX"], TEAL),
]

y = 1.32
for name, role, chips, col in layers:
    box(s1, 0.55, y, 6.60, 1.02, WHITE, LINE, 0.07, 0.9)
    box(s1, 0.55, y, 0.06, 1.02, col, None, 0.0)
    txt(s1, 0.72, y + 0.10, 2.30, 0.24, name, 9.5, True, col)
    txt(s1, 0.72, y + 0.40, 2.30, 0.52, role, 7.8, False, MUT, spacing=1.15)
    cx = 3.10
    for c in chips:
        wch = 0.088 * len(c) + 0.26
        if cx + wch > 6.95:
            cx = 3.10
            y_chip = y + 0.58
        else:
            y_chip = y + 0.26
        box(s1, cx, y_chip, wch, 0.32, SOFT, col, 0.14, 1.0)
        txt(s1, cx, y_chip + 0.06, wch, 0.22, c, 8, True, col, PP_ALIGN.CENTER)
        cx += wch + 0.10
    y += 1.10

txt(s1, 0.55, 6.86, 6.60, 0.26,
    "越靠上越「业务」，越靠下越「电气」；换层不改相邻层，是这套分层能长期使用的原因。", 7.8, False, MUT)

pic(s1, "teststand_hero.png", 7.45, 1.30, 5.30,
    "图 1　NI TestStand 产品页配图：序列编辑与执行监视（来源 ni.com）")
pic(s1, "pxi_config.png", 7.45, 4.28, 5.30,
    "图 2　PXI 机箱与系统配置界面，含 InstrumentStudio 通道视图（来源 ni.com）")
footer(s1, "图片采集自 ni.com 产品页（2026-10），仅用于内部技术评审；软件分层为本文归纳，非 NI 官方分类")

# ─────────────────────────── P2：角色表 + 执行链
s2 = prs.slides.add_slide(BLANK)
header(s2, "每个软件到底在干什么：角色与交接点",
       "把「谁负责哪一步」写清楚，才能判断哪些能力可以自建、哪些必须采购", 2)

chain = [("① 配置", "NI MAX · PXI Platform Services", TEAL),
         ("② 序列", "TestStand（Process Model）", RED),
         ("③ 并行", "TSM：site / pin map", RED),
         ("④ 测量", "LabVIEW VI", BLUE),
         ("⑤ 路由与驱动", "Switch Executive · VISA", GREEN),
         ("⑥ 归档", "SystemLink · DIAdem", PURPLE)]
for i, (a, b, col) in enumerate(chain):
    row, col_i = divmod(i, 3)
    x = 0.55 + col_i * 2.62
    yy = 1.28 + row * 0.80
    box(s2, x, yy, 2.48, 0.70, WHITE, col, 0.10, 1.3)
    txt(s2, x + 0.12, yy + 0.09, 2.24, 0.22, a, 8.5, True, col)
    txt(s2, x + 0.12, yy + 0.36, 2.24, 0.28, b, 7.5, False, MUT, spacing=1.05)

rows = [
    ("TestStand", "序列引擎：解释执行 step、跑流程模型、收结果", "②③", RED),
    ("TSM", "site 与 pin map、多工位复用、binning 与 STDF", "③", RED),
    ("LabVIEW", "测量代码本体，一个 step 对应一个 VI", "④", BLUE),
    ("Switch Executive", "开关与矩阵路径配置、共享资源互斥", "⑤", GREEN),
    ("InstrumentStudio", "仪器软面板：交互调试、参数试探、截图留证", "④", BLUE),
    ("NI-VISA / NI-488.2", "统一会话层：LAN / GPIB / USB 仪器同一套接口", "⑤", GREEN),
    ("NI-DAQmx", "传感器与模拟量采集、定时与触发", "⑤", GREEN),
    ("PXI Platform Services / MAX", "机箱与模块发现、时钟触发配置、自检", "①", TEAL),
    ("SystemLink / DIAdem", "机群与测试数据管理；离线分析与报告", "⑥", PURPLE),
]
ty = 3.10
box(s2, 0.55, ty, 7.85, 0.30, RGBColor(0xEE, 0xEC, 0xE6), None, 0.05)
for j, h in enumerate(["软件", "在 ATE 里的角色", "交接点"]):
    txt(s2, 0.68 + [0, 1.55, 6.55][j], ty + 0.05, [1.40, 4.40, 1.50][j], 0.20, h, 7.8, True, MUT)
for i, (a, b, c, col) in enumerate(rows):
    ry = ty + 0.30 + i * 0.415
    box(s2, 0.55, ry, 7.85, 0.38, WHITE if i % 2 == 0 else SOFT, LINE, 0.06, 0.7)
    box(s2, 0.55, ry, 0.04, 0.38, col, None, 0.0)
    txt(s2, 0.68, ry + 0.09, 1.30, 0.22, a, 7.8, True, INK)
    txt(s2, 2.05, ry + 0.09, 4.55, 0.22, b, 7.4, False, RGBColor(0x3A, 0x38, 0x34))
    txt(s2, 6.62, ry + 0.09, 1.55, 0.22, c, 7.8, True, col)

pic(s2, "labview_hero.png", 8.65, 1.28, 4.15,
    "图 3　LabVIEW 产品页配图：图形化程序与前面板（来源 ni.com）")
pic(s2, "x_PXIe-1081_Configuration_2_InstrumentStudio_540x323_2x.png", 8.65, 3.96, 4.15,
    "图 4　PXIe-1081 机箱与 InstrumentStudio 同屏（来源 ni.com）")

box(s2, 8.65, 6.54, 4.15, 0.52, CREAM, AMBER, 0.08, 1.2)
txt(s2, 8.78, 6.62, 3.92, 0.40,
    "买 vs 自建：②③⑥ 是通用能力，采购更划算；\n④⑤ 与夹具、被测件强绑定，通常自研。",
    7.6, False, RGBColor(0x4A, 0x3A, 0x1E), spacing=1.2)
footer(s2, "角色描述依据 NI 官方文档与产品页；「买 vs 自建」为本文判断")

prs.save(OUT)

p = Presentation(OUT)
pics = oob = ov = 0
for si, sl in enumerate(p.slides):
    boxes = []
    for sh in sl.shapes:
        if sh.shape_type == 13:
            pics += 1
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界 p{si+1}: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f}) {getattr(sh,'name','')}")
        if sh.has_text_frame and sh.text_frame.text.strip():
            boxes.append((l, t, w, h, sh.text_frame.text[:12]))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠 p{si+1}: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜图片 {pics}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
