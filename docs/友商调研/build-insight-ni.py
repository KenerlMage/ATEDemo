# -*- coding: utf-8 -*-
"""业界方案洞察 · NI：P1 硬件拓扑、P2 软件架构
页标题统一「业界方案洞察」；标题下灰底文本框作公司简介；页底小结；右侧配图。"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
IMGD = os.path.join(HERE, "ni_img")
OUT = os.path.join(HERE, "业界方案洞察-NI.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x33, 0x33, 0x33)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
RULE = RGBColor(0xC4, 0xC0, 0xB6)
FILL = RGBColor(0xF7, 0xF6, 0xF3)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def crop(src, dst, ar, focus=0.5, vfocus=0.5):
    im = Image.open(os.path.join(IMGD, src)).convert("RGB")
    w, h = im.size
    tw, th = w, w / ar
    if th > h:
        th, tw = h, h * ar
    left = max(0, min(w - tw, w * focus - tw / 2))
    top = max(0, min(h - th, h * vfocus - th / 2))
    im.crop((int(left), int(top), int(left + tw), int(top + th))).save(
        os.path.join(IMGD, dst), "JPEG", quality=90, optimize=True)


crop("pxi_config.png", "rep_pxi.jpg", 3.82, 0.52, 0.5)
crop("teststand_hero.png", "rep_ts.jpg", 2.00, 0.44, 0.5)
crop("labview_hero.png", "rep_lv.jpg", 2.00, 0.50, 0.5)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(sl, x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.2):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = ln
        _font(r, size, bold, color)
    return tb


def rect(sl, x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(sl, x, y, w, color=RULE, lw=0.012):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(sl, x, y, h, color=RED, w=0.045):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def varrow(sl, x, y, h=0.20, color=RGBColor(0xAD, 0xA9, 0x9F)):
    s = sl.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(0.13), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.text_frame.text = ""
    return s


def head(sl, sub_lead, sub_rest, page):
    txt(sl, 0.56, 0.24, 8.0, 0.38, "业界方案洞察", 21, True, INK)
    txt(sl, 8.90, 0.36, 3.90, 0.22, sub_lead, 8.6, False, DIM, PP_ALIGN.RIGHT)
    rect(sl, 0.56, 0.70, 12.20, 0.64, GRAY, None, 0)
    vrule(sl, 0.56, 0.70, 0.64, RED, 0.05)
    txt(sl, 0.74, 0.80, 11.85, 0.48, sub_rest, 7.6, False, RGBColor(0x2E, 0x2E, 0x2E), ls=1.25)
    txt(sl, 12.30, 7.14, 0.50, 0.20, page, 8, False, DIM, PP_ALIGN.RIGHT)


def band(sl, x, y, w, h, label, sub=""):
    rect(sl, x, y, w, h, WHITE, RULE, 0.75)
    vrule(sl, x, y, h, RED, 0.045)
    txt(sl, x + 0.11, y + 0.09, 0.92, 0.20, label, 8.0, True, RED)
    if sub:
        txt(sl, x + 0.11, y + 0.31, 0.92, h - 0.36, sub, 6.4, False, DIM)


def cell(sl, x, y, w, h, name, desc, i):
    rect(sl, x, y, w, h, FILL if i % 2 == 0 else WHITE, RULE, 0.75)
    txt(sl, x + 0.09, y + 0.08, w - 0.18, 0.19, name, 7.4, True, INK)
    txt(sl, x + 0.09, y + 0.29, w - 0.18, h - 0.33, desc, 6.6, False, MUT, ls=1.15)


def rhead(sl, x, y, w, text):
    txt(sl, x, y, w, 0.18, text, 8.2, True, RED)
    hline(sl, x, y + 0.23, w, RULE)


def scene(sl, x, y, w, rows, pitch=0.30):
    for i, (a, b) in enumerate(rows):
        rect(sl, x, y + i * pitch, w, pitch, WHITE if i % 2 else FILL, RULE, 0.5)
        txt(sl, x + 0.10, y + i * pitch + 0.07, w - 0.92, 0.18, a, 7.0, False, BODY)
        txt(sl, x + w - 0.72, y + i * pitch + 0.07, 0.62, 0.18, b, 7.0, True,
            RED if b == "受限" else INK, PP_ALIGN.RIGHT)


def summary(sl, text):
    rect(sl, 0.56, 6.34, 12.20, 0.58, WHITE, RULE, 0.75)
    vrule(sl, 0.56, 6.34, 0.58, RED, 0.05)
    txt(sl, 0.70, 6.41, 0.52, 0.20, "小结", 8.0, True, RED)
    txt(sl, 1.22, 6.42, 11.40, 0.44, text, 7.8, False, RGBColor(0x25, 0x25, 0x25), ls=1.2)


def foot(sl, s):
    hline(sl, 0.56, 7.02, 12.20, RULE, 0.010)
    txt(sl, 0.56, 7.09, 11.5, 0.20, s, 6.6, False, DIM)


def pic(sl, path, x, y, w, h, cap):
    sl.shapes.add_picture(os.path.join(IMGD, path), Inches(x), Inches(y),
                          width=Inches(w), height=Inches(h))
    rect(sl, x, y, w, h, None, RULE, 0.75)
    txt(sl, x, y + h + 0.03, w, 0.18, cap, 6.4, False, DIM)


ABOUT_HW = ("美国国家仪器（National Instruments，简称 NI）成立于 1976 年，总部位于美国得克萨斯州奥斯汀，"
            "长期提供模块化仪器、测试执行软件与系统集成工具；2023 年并入 Emerson，作为其测试与测量业务运营。"
            "其 PXI 模块化平台与 TestStand 测试执行软件在电子、汽车与半导体测试领域应用广泛。"
            "本页基于其公开资料，说明该方案的硬件拓扑与适用场景。")
ABOUT_SW = ("NI 的软件体系以 LabVIEW（1986 年首次发布）为基础，向上延伸出 TestStand 测试执行软件，"
            "以及面向半导体测试的 TestStand Semiconductor Module。其公开资料将测量、测试执行与数据管理划分为"
            "相对独立的产品线，各自独立发布与升级。本页说明该方案的软件分层与适用场景。")

LY, LH, LG = 1.50, 1.00, 0.26
TX, TW = 0.56, 7.30
RX, RW = 8.02, 4.74

# ═══════════════════════════════ P1 硬件拓扑
s1 = prs.slides.add_slide(BLANK)
head(s1, "第一篇　硬件拓扑", ABOUT_HW, "1 / 2")

bands = [
    ("中心与网络", "管理与数据", [
        ("工业以太网交换", "站点与中心互联\n千兆及以上"),
        ("中心服务器", "程序库与结果库\n部署与追溯"),
        ("制造执行接口", "任务令与结果回传\n按产线接口约定"),
    ]),
    ("站点控制", "工控计算机", [
        ("工业控制计算机", "运行测试执行软件\n连接机箱与外部仪器"),
        ("嵌入式控制器", "插于机箱内部\n适用于紧凑型工位"),
        ("操作终端", "启动、监视与处置\n不参与测量控制"),
    ]),
    ("仪器与开关", "测量资源", [
        ("模块化机箱", "数字、模拟、射频模块\n背板提供时钟与触发"),
        ("开关与矩阵", "通道扩展与路径切换\n共享资源的互斥管理"),
        ("外挂仪器", "以太网、总线与串口仪器\n统一会话访问"),
    ]),
    ("接口与被测", "被测对象", [
        ("夹具与针床", "机械定位与信号引出\n更换品种时整体替换"),
        ("被测件工位", "同一机箱支持多工位\n各工位并行执行"),
        ("上下料装置", "手动或自动上下料\n与测试节拍配合"),
    ]),
]

for i, (label, sub, cells) in enumerate(bands):
    y = LY + i * (LH + LG)
    band(s1, TX, y, TW, LH, label, sub)
    cw = (TW - 1.13 - 0.18) / 3
    for j, (nm, ds) in enumerate(cells):
        cell(s1, TX + 1.13 + j * (cw + 0.09), y + 0.09, cw, LH - 0.18, nm, ds, j)
    if i < 3:
        varrow(s1, TX + TW / 2 - 0.065, y + LH + 0.04, LG - 0.08)

pic(s1, "rep_pxi.jpg", RX, 1.50, RW, 1.24, "图 1　模块化机箱与仪器配置界面（ni.com 产品资料）")

rhead(s1, RX, 2.96, RW, "拓扑说明")
notes = [
    ("（一）", "单站点纵向结构：中心侧只负责程序下发与结果回收，站点侧承担全部实时测试动作。"),
    ("（二）", "多工位由机箱内并行执行实现，不通过增加工业控制计算机数量扩展工位；"
              "程序与结果经工业以太网传输，仪器控制信号经背板传输，两条通路相互独立。"),
]
ny = 3.20
for k, t in notes:
    txt(s1, RX, ny, 0.42, 0.18, k, 7.2, True, DIM)
    txt(s1, RX + 0.42, ny, RW - 0.42, 0.56, t, 7.2, False, BODY, ls=1.2)
    ny += 0.62

rhead(s1, RX, 4.50, RW, "适用场景")
scene(s1, RX, 4.76, RW, [
    ("多品种、中小批量、引脚数中等以上", "适用"),
    ("单一机箱内 4 至 16 工位并行", "适用"),
    ("数字、模拟与射频混合测量", "适用"),
    ("成本敏感的单品种大批量", "受限"),
    ("需与既有非 NI 仪器深度耦合", "受限"),
])

summary(s1, "该方案以单一站点承载多工位并行，扩展工位不依赖增加计算机数量；程序与结果、仪器控制分走两条通路，"
            "适用于中等规模、品种较多的量产线。适用性判断以功能覆盖与工程成本为依据，不含价格因素。")
foot(s1, "资料来源：NI 公开产品资料与公司信息；分层方式与适用性判断为本文归纳。")

# ═══════════════════════════════ P2 软件架构
s2 = prs.slides.add_slide(BLANK)
head(s2, "第二篇　软件架构", ABOUT_SW, "2 / 2")

layers = [
    ("应用层", "使用者视角", "操作界面　程序管理　结果查询",
     "面向操作与工艺人员，负责程序选择、运行监视与结果查阅，不直接控制仪器。"),
    ("执行层", "流程组织", "序列执行　工位管理　判定与分档",
     "按既定流程组织测试步骤，管理多工位并行，汇集测量结果并完成合格性判定与分档。"),
    ("测量层", "测量实现", "测量程序　仪器驱动　开关路径",
     "面向具体测试项，完成激励施加、信号采集与数据换算，并与仪器和开关资源对接。"),
    ("平台层", "资源与配置", "机箱与模块管理　系统配置　自检",
     "完成硬件资源发现、通道与触发配置、设备自检与固件维护，为以上各层提供统一资源视图。"),
]

for i, (label, sub, comps, desc) in enumerate(layers):
    y = LY + i * (LH + LG)
    band(s2, TX, y, TW, LH, label, sub)
    bx, bw = TX + 1.13, TW - 1.13
    rect(s2, bx, y + 0.09, bw, LH - 0.18, WHITE, RULE, 0.75)
    txt(s2, bx + 0.13, y + 0.16, bw - 0.26, 0.19, comps, 7.4, True, BLUE)
    txt(s2, bx + 0.13, y + 0.41, bw - 0.26, LH - 0.52, desc, 7.2, False, BODY, ls=1.18)
    if i < 3:
        varrow(s2, TX + TW / 2 - 0.065, y + LH + 0.04, LG - 0.08)

pic(s2, "rep_ts.jpg", RX, 1.50, 2.32, 1.16, "图 1　测试执行软件")
pic(s2, "rep_lv.jpg", RX + 2.42, 1.50, 2.32, 1.16, "图 2　测量开发环境")

rhead(s2, RX, 2.88, RW, "架构特征")
feats = [
    ("（一）", "层次职责单一：应用层不涉及测量，测量层不涉及流程，任一层调整不影响相邻层。"),
    ("（二）", "流程与测量分离、结果出口统一：换被测对象主要改测量层与流程配置，换仪器主要改测量层；"
              "各层结果按统一格式归档，供中心侧汇总与追溯。"),
]
ny = 3.16
for k, t in feats:
    txt(s2, RX, ny, 0.42, 0.18, k, 7.2, True, DIM)
    txt(s2, RX + 0.42, ny, RW - 0.42, 0.58, t, 7.2, False, BODY, ls=1.2)
    ny += 0.64

rhead(s2, RX, 4.50, RW, "适用场景")
scene(s2, RX, 4.76, RW, [
    ("流程稳定、需长期维护的量产线", "适用"),
    ("测试项由工艺定义、软件仅作组织", "适用"),
    ("多品种小批量、换型频繁", "适用"),
    ("频繁变更的研发验证阶段", "受限"),
    ("仅需轻量单工位执行的小型工位", "受限"),
])

summary(s2, "软件按应用、执行、测量、平台四层划分，流程与测量分离使换型与换仪器的影响范围可控；"
            "该结构适合流程稳定、需长期维护的量产场景，研发验证阶段因配置与授权调整周期较长而相对受限。")
foot(s2, "资料来源：NI 公开产品资料与公司信息；四层划分与适用性判断为本文归纳，非厂商官方分类。")

prs.save(OUT)

p = Presentation(OUT)
pics = oob = ov = 0
for si, sl in enumerate(p.slides):
    bx = []
    for sh in sl.shapes:
        if sh.shape_type == 13:
            pics += 1
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界 p{si+1}: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bx.append((l, t, w, h, sh.text_frame.text[:12]))
    for i in range(len(bx)):
        for j in range(i + 1, len(bx)):
            a, b = bx[i], bx[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠 p{si+1}: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜图片 {pics}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
