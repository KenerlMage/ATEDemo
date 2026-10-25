# -*- coding: utf-8 -*-
"""业畍方案洞察 · 一页归并版 v2：产品两列换行排布 + 蓝色关键词 + 右侧产品图片拓扑（半导体联动）"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
IMGD = os.path.join(HERE, "ni_img")
OUT = os.path.join(HERE, "业界方案洞察-NI-一页归并版-v2.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
BLUE2 = RGBColor(0x1B, 0x5E, 0xA8)
RULE = RGBColor(0xBF, 0xBB, 0xB1)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR = RGBColor(0xEE, 0xEC, 0xE6)
WIRE = RGBColor(0x6E, 0x6A, 0x62)


def crop(src, dst, ar=1.50, focus=0.5, q=86):
    im = Image.open(os.path.join(IMGD, src)).convert("RGB")
    w, h = im.size
    tw, th = w, w / ar
    if th > h:
        th, tw = h, h * ar
    left = max(0, min(w - tw, w * focus - tw / 2))
    top = max(0, (h - th) / 2)
    out = im.crop((int(left), int(top), int(left + tw), int(top + th)))
    if out.size[0] > 620:
        out = out.resize((620, int(620 / ar)), Image.LANCZOS)
    out.save(os.path.join(IMGD, dst), "JPEG", quality=q, optimize=True)


crop("labview_hero.png", "n2_lv.jpg", 1.50, 0.5)
crop("teststand_table.png", "n2_tsm.jpg", 1.50, 0.5)
crop("teststand_hero.png", "n2_ts.jpg", 1.50, 0.5)
crop("node_chassis.jpg", "n2_pxi.jpg", 1.50, 0.45)
crop("x_PXIe-1081_Configuration_2_InstrumentStudio_540x323_2x.png", "n2_studio.jpg", 1.50, 0.5)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def para(tf, first, segs, size, bold, color, align, ls, italic=False):
    """segs: [(text, color_or_None, bold_or_None)]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = ls
    for t, c, b in segs:
        r = p.add_run()
        r.text = t
        _font(r, size, bold if b is None else b, color if c is None else c, italic)
    return p


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14, italic=False):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        para(tf, i == 0, [(ln, None, None)], size, bold, color, align, ls, italic)
    return tb


def txt_hi(x, y, w, h, text, size=6.2, color=BODY, hi=BLUE2, hi_bold=True,
           ls=1.18, align=PP_ALIGN.LEFT):
    """用 ** 标记核心内容，渲染为蓝色（加粗）"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    parts = str(text).split("**")
    segs = []
    for i, seg in enumerate(parts):
        if not seg:
            continue
        on = (i % 2 == 1)
        segs.append((seg, hi if on else None, hi_bold if on else None))
    para(tf, True, segs, size, False, color, align, ls)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.012):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(x, y, h, color=RED, w=0.045):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def wire(pts, color=WIRE, lw=1.0, arrow=True):
    for i in range(len(pts) - 1):
        x1, y1 = pts[i]
        x2, y2 = pts[i + 1]
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(lw)
        ln = c.line._get_or_add_ln()
        if arrow and i == len(pts) - 2:
            t = etree.SubElement(ln, qn("a:tailEnd"))
            t.set("type", "triangle")
            t.set("w", "med")
            t.set("len", "med")


def secno(x, y, no, name, rule_w=1.66):
    txt(x, y - 0.04, 0.34, 0.30, no, 13.5, True, RED)
    txt(x + 0.40, y + 0.02, 1.90, 0.18, name, 8.0, True, RED)
    hline(x + 0.40, y + 0.24, rule_w, RULE, 0.010)


def picnode(x, y, w, h, img, label, sub=None, accent=BLUE):
    rect(x, y, w, h, WHITE, accent, 0.9)
    vrule(x, y, h, accent, 0.04)
    iw, ih = 0.56, 0.56 / 1.50
    iy = y + (h - ih) / 2
    sl.shapes.add_picture(os.path.join(IMGD, img), Inches(x + 0.07), Inches(iy),
                          width=Inches(iw), height=Inches(ih))
    if sub:
        txt(x + 0.70, y + 0.055, w - 0.78, 0.16, label, 6.2, True, INK)
        txt(x + 0.70, y + 0.235, w - 0.78, 0.16, sub, 5.5, False, MUT)
    else:
        txt(x + 0.70, y, w - 0.78, h, label, 6.2, True, INK, anchor=MSO_ANCHOR.MIDDLE)


def txtnode(x, y, w, h, label, sub=None, accent=RULE, fill=FILL2):
    rect(x, y, w, h, fill, accent, 0.75)
    vrule(x, y, h, accent if accent != RULE else RED, 0.04)
    if sub:
        txt(x + 0.12, y + 0.055, w - 0.20, 0.16, label, 6.2, True, INK)
        txt(x + 0.12, y + 0.235, w - 0.20, 0.16, sub, 5.5, False, MUT)
    else:
        txt(x + 0.12, y, w - 0.20, h, label, 6.2, True, INK, anchor=MSO_ANCHOR.MIDDLE)


# ── 页头
txt(0.56, 0.20, 7.60, 0.36, "业界方案洞察", 20, True, INK)
txt(8.40, 0.31, 4.36, 0.22, "NI：公司 · 产品与特性 · 核心思想", 8.2, False, DIM, PP_ALIGN.RIGHT)

# ── 一、公司
rect(0.56, 0.58, 12.20, 0.66, GRAY, None, 0)
vrule(0.56, 0.58, 0.66, RED, 0.05)
txt(0.70, 0.64, 0.30, 0.26, "一", 12.5, True, RED)
txt(0.70, 0.92, 0.60, 0.18, "公司", 6.8, True, RED)
txt_hi(1.40, 0.66, 11.24, 0.52,
       "美国国家仪器（NI）1976 年成立于奥斯汀，2023 年并入 Emerson 作为其测试与测量业务；以**测量与控制软件平台**为核心，"
       "配合**模块化硬件**，服务**半导体**、汽车、航空航天与国防、能源及院校科研。1986 年发布 LabVIEW，确立**虚拟仪器**路线："
       "功能由软件定义，硬件只提供通用的采集与激励能力。",
       6.9, RGBColor(0x2B, 0x2B, 0x2B), BLUE2, True, 1.32)

# ── 二、产品与产品特性（两列 × 三行）
secno(0.56, 1.36, "二", "产品与产品特性", 1.90)

prods = [
    ("TestStand", "测试执行与管理",
     "序列化组织测试步骤，支持**并行与多工位**；结果入库与报表；可调用 **LabVIEW / C / Python** 模块。"),
    ("LabVIEW", "图形化开发环境",
     "数据流图形编程；**仪器连接与驱动调用**；内建界面控件；**测试库与工具包**缩短开发调试周期。"),
    ("TSM", "半导体测试扩展",
     "TestStand 之上增加**工位与引脚映射**、多工位并行；对接**探针台 / 分选机**，结果以 **STDF** 输出。"),
    ("PXI 平台", "模块化硬件",
     "机箱 + 控制器 + 模块组合；**背板触发与时钟同步**；**换模块即换仪器构成**，无需专用整机。"),
    ("NI-VISA / IVI", "仪器抽象层",
     "统一会话访问 **GPIB / LXI / USB / PXI**；**类驱动实现同类仪器互换**，上层调用不随硬件变更。"),
    ("InstrumentStudio", "仪器软面板",
     "以**台式仪器面板**操作模块化仪器；面板**可保存可复用**，便于调试与人工验证。"),
]
COLW, CW2, RGAP = 3.20, 0.90, 0.12
COLS = [0.56, 0.56 + COLW + RGAP]
for i, (nm, cat, fe) in enumerate(prods):
    cx = COLS[i % 2]
    cy = 1.64 + (i // 2) * (CW2 + 0.06)
    txt(cx, cy + 0.02, COLW, 0.17, nm, 7.2, True, INK)
    txt(cx, cy + 0.185, COLW, 0.15, cat, 5.6, False, MUT)
    txt_hi(cx, cy + 0.36, COLW, 0.48, fe, 6.1, BODY, BLUE2, True, 1.20)
    hline(cx, cy + CW2, COLW, RGBColor(0xDD, 0xD9, 0xD0), 0.008)

# 列间竖线
vline_x = 0.56 + COLW + RGAP / 2
rect(vline_x, 1.66, 0.006, 2.72, RGBColor(0xDD, 0xD9, 0xD0), None, 0)

# ── 右侧：产品图片拓扑
RX, RW = 7.44, 5.32
txt(RX, 1.34, 3.24, 0.18, "产品拓扑：半导体产线中的联动", 7.4, True, RED)
txt(RX + 3.30, 1.345, RW - 3.30, 0.16, "图为 NI 公开产品图，箭头为控制与数据流向", 5.2, False, DIM, PP_ALIGN.RIGHT)
hline(RX, 1.56, RW, RULE, 0.010)

PIC_W, PIC_H = 2.62, 0.48
TX_W = 2.54
L1, L2, L3, L4, L5 = 1.64, 2.225, 2.81, 3.395, 3.98

picnode(RX, L1, PIC_W, PIC_H, "n2_ts.jpg", "TestStand", "执行与工位调度")
picnode(RX + PIC_W + 0.14, L1, TX_W, PIC_H, "n2_tsm.jpg", "TSM", "引脚映射 · 多工位")
wire([(RX + PIC_W, L1 + PIC_H / 2), (RX + PIC_W + 0.14, L1 + PIC_H / 2)])

picnode(RX, L2, PIC_W, PIC_H, "n2_lv.jpg", "LabVIEW", "测量程序 · 面板")
txtnode(RX + PIC_W + 0.14, L2, TX_W, PIC_H, "MES / STDF", "结果入库与产线回流", RED)

SPINE = RX + PIC_W / 2
MID = RX + PIC_W + 0.14 + TX_W / 2
wire([(SPINE, L1 + PIC_H), (SPINE, L2)])
wire([(SPINE, L1 + PIC_H), (SPINE, L2 - 0.055), (MID, L2 - 0.055), (MID, L2)])

txtnode(RX, L3, RW, 0.44, "NI-VISA / IVI：统一会话与类驱动，屏蔽 GPIB / LXI / USB / PXI 差异")
wire([(SPINE, L2 + PIC_H), (SPINE, L3)])

picnode(RX, L4, PIC_W, PIC_H, "n2_pxi.jpg", "PXI 机箱与模块", "背板触发 · 时钟同步")
picnode(RX + PIC_W + 0.14, L4, TX_W, PIC_H, "n2_studio.jpg", "InstrumentStudio", "软面板调试与验证")
wire([(SPINE, L3 + 0.44), (SPINE, L4)])
wire([(SPINE, L3 + 0.44), (SPINE, L4 - 0.055), (MID, L4 - 0.055), (MID, L4)])

txtnode(RX, L5, 3.30, 0.48, "半导体测试台：探针台 / 分选机", "与 TSM 交换工位与分选信号")
txtnode(RX + 3.42, L5, 1.90, 0.48, "DUT：晶圆 / 芯片")
wire([(SPINE, L4 + PIC_H), (SPINE, L5)])
wire([(RX + 3.30, L5 + 0.24), (RX + 3.42, L5 + 0.24)])

# ── 三、核心思想
secno(0.56, 4.66, "三", "核心思想", 1.90)
ideas = [
    ("软件定义仪器", "功能由软件配置决定，硬件只提供通用采集与激励；同一机箱加载不同程序即成为不同仪器。"),
    ("模块化与复用", "机箱、模块与背板同步构成复用基础；多种测量需求以组合覆盖，而非逐项定制整机。"),
    ("开放标准与生态", "PXI、VISA/IVI、STDF 等公开标准，使第三方模块、探针台与分选机可接入。"),
    ("描述与执行分离", "TestStand 负责编排，LabVIEW 与代码模块负责执行，驱动负责适配仪器，三者互不越界。"),
]
bx, bw, gap = 0.56, 2.95, 0.133
for i, (ti, de) in enumerate(ideas):
    x = bx + i * (bw + gap)
    rect(x, 4.98, bw, 1.22, WHITE, RULE, 0.75)
    vrule(x, 4.98, 1.22, BLUE, 0.04)
    txt(x + 0.13, 5.05, 0.42, 0.26, "0" + str(i + 1), 11.5, True, RGBColor(0xC9, 0xC4, 0xB8))
    txt(x + 0.13, 5.32, bw - 0.26, 0.18, ti, 7.4, True, INK)
    txt_hi(x + 0.13, 5.55, bw - 0.26, 0.58, de, 6.1, BODY, BLUE2, False, 1.14)

# ── 小结
hline(0.56, 6.44, 12.20, RULE, 0.010)
txt(0.56, 6.50, 0.52, 0.20, "小结", 8.0, True, RED)
txt_hi(1.12, 6.50, 11.60, 0.44,
       "半导体场景下分工明确：**TestStand 编排工位与流程**，**TSM 承担引脚映射与多工位并行**，"
       "**驱动层对接探针台与分选机**，结果以 **STDF 回流 MES**；对应本方案，即描述与执行分离、仪器访问抽象、仿真与实机共用同一描述。",
       7.4, RGBColor(0x25, 0x25, 0x25), BLUE2, True, 1.22)

txt(0.56, 7.09, 11.4, 0.20,
    "资料来源：NI 公开公司与产品资料（TestStand、LabVIEW、TestStand Semiconductor Module、PXI、NI-VISA、InstrumentStudio）；"
    "产品特性为公开资料归纳，拓扑与半导体联动路径为本文判断。", 6.0, False, DIM)
txt(12.30, 7.14, 0.50, 0.20, "1 / 1", 8, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
pics = oob = ov = 0
for si, s in enumerate(p.slides):
    bxs = []
    for sh in s.shapes:
        if sh.shape_type == 13:
            pics += 1
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bxs.append((l, t, w, h, sh.text_frame.text[:10]))
    for i in range(len(bxs)):
        for j in range(i + 1, len(bxs)):
            a, b = bxs[i], bxs[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜图片 {pics}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
