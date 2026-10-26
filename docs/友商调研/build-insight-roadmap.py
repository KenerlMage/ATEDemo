# -*- coding: utf-8 -*-
"""在《业界方案洞察-两种思路》后追加一页：NI 的 ATE 路标与价值"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "业界方案洞察-两种思路.pptx")
OUT = os.path.join(HERE, "业界方案洞察-两种思路-含路标.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
BLUE2 = RGBColor(0x1B, 0x5E, 0xA8)
TEAL = RGBColor(0x0F, 0x5E, 0x5A)
RULE = RGBColor(0xAD, 0xA8, 0x9C)
GRID = RGBColor(0xC6, 0xC1, 0xB5)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(SRC)

# 第一页页码 1/1 → 1/2
for sh in prs.slides[0].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "1 / 1":
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.text = "1 / 2"

sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def para(tf, first, segs, size, bold, color, align, ls):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = ls
    for t, c, b in segs:
        r = p.add_run()
        r.text = t
        _font(r, size, bold if b is None else b, color if c is None else c)
    return p


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        para(tf, i == 0, [(ln, None, None)], size, bold, color, align, ls)
    return tb


def txt_hi(x, y, w, h, text, size=6.2, color=BODY, hi=BLUE2, hi_bold=True,
           ls=1.18, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    segs = []
    for i, seg in enumerate(str(text).split("**")):
        if not seg:
            continue
        on = (i % 2 == 1)
        segs.append((seg, hi if on else None, hi_bold if on else None))
    para(tf, True, segs, size, False, color, align, ls)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.010):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def shead(x, y, w, text, color=RED, size=8.2):
    txt(x, y, w, 0.20, text, size, True, color)
    hline(x, y + 0.225, w, GRID, 0.010)


def dot(cx, cy, d=0.135, color=RED):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                            Inches(d), Inches(d))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.color.rgb = WHITE
    s.line.width = Pt(1.2)


# ── 页头
txt(0.56, 0.16, 8.40, 0.34, "NI 的 ATE 路标与价值", 19, True, INK)
txt(9.20, 0.24, 3.56, 0.22, "四个阶段 · 五条价值 · 三点取用", 8.0, False, DIM, PP_ALIGN.RIGHT)
hline(0.56, 0.54, 12.20, RULE, 0.014)

# ── 一、技术路标
shead(0.56, 0.64, 12.20, "技术路标：从仪器控制到平台化 ATE")

AXY = 1.14
stages = [
    ("1976—1995", "仪器控制与总线",
     ["首款 GPIB 接口板，解决计算机与仪器互联", "IEEE 488 总线随仪器普及", "驱动与 VISA 抽象层成形"]),
    ("1996—2010", "模块化与标准化",
     ["1997 年推出 PXI 标准，随后组建 PXI 系统联盟", "TestStand 以序列方式管理测试执行", "IVI 类驱动推动仪器可互换"]),
    ("2011—2019", "平台化 ATE",
     ["基于 PXI 与 TestStand 形成半导体测试系统", "TSM 提供多工位、引脚与通道映射", "InstrumentStudio 软面板降低调试门槛"]),
    ("2020—", "系统级与生态",
     ["从芯片测试扩展到板级与系统级测试", "面向电池、ADAS、通信的测试方案", "2023 年并入 Emerson，成为测试与测量业务"]),
]
SW4, SGAP = 2.90, 0.20
X0 = 0.56
hline(X0, AXY, 12.20, GRID, 0.012)
for i, (yr, nm, items) in enumerate(stages):
    x = X0 + i * (SW4 + SGAP)
    dot(x + 0.06, AXY, 0.135, RED if i == 0 else BLUE)
    txt(x, AXY + 0.14, SW4, 0.18, yr, 7.4, True, RED if i == 0 else BLUE)
    txt(x, AXY + 0.35, SW4, 0.20, nm, 7.8, True, INK)
    txt(x, AXY + 0.60, SW4, 0.72,
        "· " + "\n· ".join(items), 5.8, False, MUT, ls=1.26)

# ── 二、价值
shead(0.56, 2.72, 12.20, "价值：NI 为什么能以平台方式做 ATE")

vals = [
    ("01", "一套硬件覆盖多需求", "PXI 模块按测量类型组合，减少专用整机数量。"),
    ("02", "改型依靠软件", "测量能力由程序与模块决定，改型周期短。"),
    ("03", "仪器可替换", "VISA / IVI 与标准总线，避免被单一供应商锁定。"),
    ("04", "分层可分工", "执行、测量、驱动分离，团队按角色并行开发。"),
    ("05", "生态与人才", "工具链成熟、资料与工程师充足，维护门槛低。"),
]
VW, VGAP = 2.32, 0.15
for i, (no, ti, de) in enumerate(vals):
    x = 0.56 + i * (VW + VGAP)
    rect(x, 2.98, VW, 1.42, WHITE, GRID, 0.8)
    rect(x, 2.98, VW, 0.045, BLUE, None, 0)
    txt(x + 0.13, 3.08, 0.34, 0.20, no, 8.6, True, RGBColor(0xC9, 0xC4, 0xB8))
    txt(x + 0.13, 3.32, VW - 0.26, 0.34, ti, 7.0, True, INK, ls=1.12)
    txt(x + 0.13, 3.74, VW - 0.26, 0.60, de, 6.0, False, BODY, ls=1.20)

# ── 三、对我们的取用
shead(0.56, 4.58, 12.20, "对我们的取用", color=RED)

takes = [
    ("可复用", "分层与标准接口", "执行、测量、驱动三层分工，接口走 VISA / IVI 与 PXI 一类公开标准，仪器可换。", BLUE2),
    ("可复用", "仿真先行", "先用仿真设备跑通流程与结果通道，再接入实机，缩短联调期。", BLUE2),
    ("需评估", "适用边界", "该路线的价值前提是多品种小批量；单品种大批量仍应对比 per-pin 专机方案的产出与成本。", RED),
]
for i, (tag, ti, de, c) in enumerate(takes):
    y = 4.84 + i * 0.50
    rect(0.56, y, 12.20, 0.44, FILL2 if i < 2 else RGBColor(0xFA, 0xF5, 0xF5), None, 0)
    rect(0.56, y, 0.035, 0.44, c, None, 0)
    txt(0.74, y + 0.13, 0.62, 0.18, tag, 6.3, True, c)
    txt(1.42, y + 0.13, 1.70, 0.18, ti, 6.8, True, INK)
    txt(3.20, y + 0.13, 9.40, 0.20, de, 6.2, False, BODY)

# ── 判断
hline(0.56, 6.44, 12.20, RULE, 0.012)
txt(0.56, 6.50, 0.52, 0.20, "判断", 8.0, True, RED)
txt_hi(1.12, 6.49, 11.60, 0.44,
       "NI 的路标是一条**把测量能力从专用硬件搬到软件与模块上**的路径；它的价值来自**可组合、可替换、可分工**，"
       "代价是在单品种大批量的极限性能上让位于 per-pin 专机。取用时应取其**分层与标准**，而非取其全部形态。",
       7.2, RGBColor(0x25, 0x25, 0x25), BLUE2, True, 1.20)

txt(0.56, 7.10, 11.4, 0.20,
    "资料来源：NI 公开公司沿革与产品资料（GPIB、PXI、TestStand、TSM、InstrumentStudio、Emerson 并购）；"
    "年份以公开发布时间为准，价值归纳与取用判断为本文观点。", 5.9, False, DIM)
txt(12.30, 7.15, 0.50, 0.20, "2 / 2", 7.6, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
oob = ov = 0
for si, s in enumerate(p.slides):
    bxs = []
    for sh in s.shapes:
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界 p{si+1}: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bxs.append((l, t, w, h, sh.text_frame.text[:10]))
    for i in range(len(bxs)):
        for j in range(i + 1, len(bxs)):
            a, b = bxs[i], bxs[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠 p{si+1}: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
