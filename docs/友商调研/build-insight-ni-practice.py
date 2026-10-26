# -*- coding: utf-8 -*-
"""友商调研 · NI 洞察重写：灰底结论 + 中线分栏（左 拓扑与产品定义 / 右 软件拓扑与设备虚拟化）"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "业界方案洞察-NI-洞察重写.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
BLUE2 = RGBColor(0x1B, 0x5E, 0xA8)
TEAL = RGBColor(0x0F, 0x5E, 0x5A)
RULE = RGBColor(0xAD, 0xA8, 0x9C)
GRID = RGBColor(0xC6, 0xC1, 0xB5)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR = RGBColor(0xEE, 0xEC, 0xE6)
WIRE = RGBColor(0x3C, 0x3A, 0x36)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = ln
        _font(r, size, bold, color)
    return tb


def hi(x, y, w, h, text, size=6.3, color=BODY, accent=BLUE2, bold=True,
       ls=1.26, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = ls
    for i, seg in enumerate(str(text).split("**")):
        if not seg:
            continue
        on = (i % 2 == 1)
        r = p.add_run()
        r.text = seg
        _font(r, size, bold if on else False, accent if on else color)
    return tb


def paras(x, y, w, h, items, size=6.3, color=BODY, accent=BLUE2, bold=True, ls=1.30, gap=4):
    """items: 段落文本列表，段间留空"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = ls
        p.space_after = Pt(gap)
        for i, seg in enumerate(str(it).split("**")):
            if not seg:
                continue
            on = (i % 2 == 1)
            r = p.add_run()
            r.text = seg
            _font(r, size, bold if on else False, accent if on else color)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.010):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(x, y, h, color=RED, w=0.040):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def arrow(x1, y1, x2, y2, color=WIRE, lw=0.9):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    t = etree.SubElement(ln, qn("a:tailEnd"))
    t.set("type", "triangle")
    t.set("w", "sm")
    t.set("len", "sm")
    return c


def head(x, y, w, text, color=RED, size=8.4):
    txt(x, y, w, 0.20, text, size, True, color)
    hline(x, y + 0.23, w, GRID, 0.010)


# ══════════ 页头 ══════════
txt(0.56, 0.16, 8.30, 0.34, "业界方案洞察：NI 在 ATE 领域的优秀实践", 19, True, INK)
txt(8.90, 0.24, 3.86, 0.22, "软件拓扑 · 设备虚拟化", 8.4, False, DIM, PP_ALIGN.RIGHT)
hline(0.56, 0.54, 12.20, RULE, 0.014)

# ══════════ 洞察结论（灰底）══════════
rect(0.56, 0.64, 12.20, 1.06, GRAY, None, 0)
vrule(0.56, 0.64, 1.06, RED, 0.048)
hi(0.76, 0.70, 11.86, 0.20,
   "洞察结论：NI 的 ATE 能力不由单台仪器决定，而由**软件拓扑**与**设备虚拟化**两项设计决定。",
   7.2, RGBColor(0x22, 0x22, 0x22), BLUE2, True, 1.20)
concl = [
    "① **软件拓扑**：应用、执行、测量、访问、硬件五层单向依赖，同层可替换——因此更换仪器不必改写测试程序。",
    "② **设备虚拟化**：仪器以通道、量程、采样率与触发方式被描述，可被仿真、可被类驱动映射——因此无实机亦可联调、跨机型只换配置。",
    "③ 两者共同的作用：把**机型切换成本**与**联调周期**从硬件侧移到软件侧，这是 NI 在多品种小批量产线上的实践基础。",
]
for i, t in enumerate(concl):
    hi(0.78, 0.95 + i * 0.245, 11.84, 0.22, t, 6.8, RGBColor(0x30, 0x30, 0x30), BLUE2, True, 1.18)

# ══════════ 黑色细中线分两栏 ══════════
CX = 6.66
rect(CX, 1.80, 0.012, 5.06, RGBColor(0x1A, 0x1A, 0x1A), None, 0)

# ── 左：拓扑与产品定义
LX, LW = 0.56, 5.80
head(LX, 1.78, LW, "NI 的系统拓扑（自下而上）")

rows = [
    ("业务与数据", "MES · 数据库 · STDF 报表", FILL2),
    ("测试执行", "TestStand ＋ TSM（工位与引脚映射）", WHITE),
    ("测量开发", "LabVIEW 测试程序", WHITE),
    ("仪器访问", "NI-VISA 会话层 ＋ IVI 类驱动", FILL2),
    ("硬件承载", "PXI 机箱与模块：DAQ / DMM / SMU / 开关 / 示波器", WHITE),
    ("被测对象", "探针台 / 分选机 / DUT / 针床", RGBColor(0xF7, 0xF4, 0xEE)),
]
BX, BW = 1.52, 3.30
for i, (role, prod, fill) in enumerate(rows):
    y = 2.02 + i * 0.50
    txt(LX, y + 0.075, 0.92, 0.18, role, 5.9, True, MUT)
    rect(BX, y, BW, 0.36, fill, BLUE if i in (1, 3) else GRID, 0.8)
    txt(BX + 0.10, y + 0.095, BW - 0.18, 0.18, prod, 5.8, False, INK)
    if i < 5:
        arrow(BX + BW / 2, y + 0.37, BX + BW / 2, y + 0.485, WIRE, 0.8)
txt(BX + BW + 0.06, 2.02 + 0.10, 0.90, 0.18, "↑ 单向依赖", 5.2, False, DIM)

# 产品定义
head(LX, 5.10, LW, "产品定义：每个组件解决什么问题", color=BLUE)
prd = [
    ("TestStand", "测试执行：序列、并行、多工位、结果与报告"),
    ("TSM", "半导体语义：工位、引脚、通道映射，对接探针台"),
    ("LabVIEW", "测量程序开发：图形化调用驱动，形成测试逻辑"),
    ("VISA / IVI", "访问抽象：统一会话，类驱动使仪器可互换"),
    ("PXI", "物理承载：机箱与模块提供通道、同步与触发"),
]
ry = 5.36
for i, (nm, de) in enumerate(prd):
    rect(LX, ry, 1.16, 0.235, FILL2, None, 0)
    txt(LX + 0.08, ry + 0.045, 1.06, 0.16, nm, 5.6, True, BLUE)
    txt(LX + 1.26, ry + 0.045, 4.52, 0.16, de, 5.7, False, BODY)
    ry += 0.275

# ── 右：两种思想
RX, RW = 6.90, 5.86

head(RX, 1.78, RW, "NI 的软件拓扑思想")
paras(RX, 2.06, RW, 1.20, [
    "**分层且单向依赖**：界面只调用执行引擎，执行引擎只调用测量程序，测量程序只经访问层触达硬件；任一层替换，其上各层语义与程序不变。",
    "**接口标准化**：会话统一走 VISA，仪器类别统一走 IVI，同一程序可运行于不同厂商、不同形态的仪器；标准接口把选型与开发解耦。",
    "**语义上移**：TSM 把工位、引脚、通道等半导体概念置于执行层之上，程序只表达测什么、判什么，不表达接在哪一路。",
    "**结果外置**：结果与量测数据以统一格式（如 STDF）流向 MES 与分析侧，测试程序不承担数据格式职责。",
], size=7.0, color=BODY, accent=BLUE2, ls=1.34, gap=3)

head(RX, 3.40, RW, "NI 的设备虚拟化思想")
paras(RX, 3.68, RW, 1.40, [
    "**以能力描述硬件**：仪器抽象为通道数、量程、采样率与触发方式的集合，程序按能力编写，不按型号编写。",
    "**以仿真先行**：以仿真设备替代实机，先跑通流程、判定与数据通道；实机到位只替换会话，程序不动。",
    "**以类驱动映射**：同一逻辑仪器映射到不同厂商实现，配置变化不触发代码变更，减少重新验证的工作量。",
    "**以配置跨形态**：同一程序可运行于 PXI 模块与台式仪器，机型切换由配置完成，而非重写程序。",
], size=7.0, color=BODY, accent=BLUE2, ls=1.34, gap=3)

# 取用
rect(RX, 5.06, RW, 1.52, FILL2, None, 0)
vrule(RX, 5.06, 1.52, TEAL, 0.042)
txt(RX + 0.18, 5.15, RW - 0.32, 0.18, "对我们的取用：顺序不可颠倒", 7.0, True, TEAL)
takes = [
    ("一", "先立访问层", "统一会话与驱动映射，仪器可换、程序不改。"),
    ("二", "再立仿真层", "无实机即可联调，联调期与硬件到货解耦。"),
    ("三", "后谈平台化", "缺此两层，分层只停在文档上，切换成本仍落回硬件与人力。"),
]
for i, (no, ti, de) in enumerate(takes):
    y = 5.40 + i * 0.37
    txt(RX + 0.18, y, 0.22, 0.16, no, 6.4, True, TEAL)
    txt(RX + 0.44, y, 1.10, 0.16, ti, 6.4, True, INK)
    txt(RX + 1.60, y, RW - 1.76, 0.32, de, 6.2, False, BODY, ls=1.16)

# 页脚
txt(0.56, 7.02, 10.9, 0.20,
    "资料来源：NI 公开产品资料与技术文档（TestStand、LabVIEW、TSM、NI-VISA、IVI、PXI）；架构归纳、洞察结论与取用判断为本文观点。",
    5.8, False, DIM)
txt(12.28, 7.05, 0.52, 0.20, "1 / 1", 7.4, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
oob = ov = 0
for si, s in enumerate(p.slides):
    bxs = []
    for sh in s.shapes:
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bxs.append((l, t, w, h, sh.text_frame.text[:12]))
    for i in range(len(bxs)):
        for j in range(i + 1, len(bxs)):
            a, b = bxs[i], bxs[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
