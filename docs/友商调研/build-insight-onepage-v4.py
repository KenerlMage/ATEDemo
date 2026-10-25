# -*- coding: utf-8 -*-
"""业界方案洞察 · v4：恢复灰底公司介绍；DSL 不再单列，并入核心思想第 4 条"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
IMGD = os.path.join(HERE, "ni_img")
OUT = os.path.join(HERE, "业界方案洞察-NI-一页归并版-v4.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
BLUE2 = RGBColor(0x1B, 0x5E, 0xA8)
RULE = RGBColor(0xAD, 0xA8, 0x9C)
GRID = RGBColor(0xC6, 0xC1, 0xB5)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WIRE = RGBColor(0x6E, 0x6A, 0x62)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def para(tf, first, segs, size, bold, color, align, ls):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = ls
    for t, c, b in segs:
        r = p.add_run()
        r.text = t
        _font(r, size, bold if b is None else b, color if c is None else c)
    return p


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        para(tf, i == 0, [(ln, None, None)], size, bold, color, align, ls)
    return tb


def txt_hi(x, y, w, h, text, size=6.2, color=BODY, hi=BLUE2, hi_bold=True,
           ls=1.18, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    segs = []
    for i, seg in enumerate(str(text).split("**")):
        if not seg:
            continue
        on = (i % 2 == 1)
        segs.append((seg, hi if on else None, hi_bold if on else None))
    para(tf, True, segs, size, False, color, align, ls)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.010):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(x, y, h, color=RED, w=0.040):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def wire(pts, color=WIRE, lw=1.0, arrow=True):
    for i in range(len(pts) - 1):
        x1, y1 = pts[i]
        x2, y2 = pts[i + 1]
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(lw)
        ln = c.line._get_or_add_ln()
        if arrow and i == len(pts) - 2:
            t = etree.SubElement(ln, qn("a:tailEnd"))
            t.set("type", "triangle")
            t.set("w", "med")
            t.set("len", "med")


def shead(x, y, w, text):
    txt(x, y, w, 0.20, text, 8.8, True, RED)
    hline(x, y + 0.235, w, RULE, 0.012)


def picnode(x, y, w, h, img, label, sub=None, accent=BLUE):
    rect(x, y, w, h, WHITE, accent, 0.9)
    vrule(x, y, h, accent, 0.04)
    iw, ih = 0.56, 0.56 / 1.50
    sl.shapes.add_picture(os.path.join(IMGD, img), Inches(x + 0.07), Inches(y + (h - ih) / 2),
                          width=Inches(iw), height=Inches(ih))
    txt(x + 0.70, y + 0.075, w - 0.78, 0.16, label, 6.2, True, INK)
    txt(x + 0.70, y + 0.255, w - 0.78, 0.16, sub, 5.5, False, MUT)


def txtnode(x, y, w, h, label, sub=None, accent=RED, fill=FILL2):
    rect(x, y, w, h, fill, RULE, 0.7)
    vrule(x, y, h, accent, 0.04)
    if sub:
        txt(x + 0.12, y + 0.075, w - 0.20, 0.16, label, 6.2, True, INK)
        txt(x + 0.12, y + 0.255, w - 0.20, 0.16, sub, 5.5, False, MUT)
    else:
        txt(x + 0.12, y, w - 0.20, h, label, 6.2, True, INK, anchor=MSO_ANCHOR.MIDDLE)


# ── 页头
txt(0.56, 0.16, 7.60, 0.36, "业界方案洞察", 20, True, INK)
txt(8.40, 0.27, 4.36, 0.22, "NI：产品与特性 · 拓扑与核心思想", 8.2, False, DIM, PP_ALIGN.RIGHT)

# ── 公司介绍（灰底）
rect(0.56, 0.58, 12.20, 0.68, GRAY, None, 0)
vrule(0.56, 0.58, 0.68, RED, 0.045)
txt_hi(0.74, 0.67, 11.90, 0.54,
       "美国国家仪器（NI）1976 年成立于美国得克萨斯州奥斯汀，2023 年并入 Emerson 作为其测试与测量业务运营；"
       "以**测量与控制软件平台**为核心，配合**模块化硬件**，服务**半导体**、汽车、航空航天与国防、能源及院校科研的验证与生产测试；"
       "1986 年发布 LabVIEW，由此确立**虚拟仪器**路线。",
       6.9, RGBColor(0x2B, 0x2B, 0x2B), BLUE2, True, 1.30)

# ── 左：产品与产品特性
shead(0.56, 1.36, 3.10, "产品与产品特性")

GX, GY, GW, GH = 0.56, 1.66, 6.58, 2.78
COLX = [0.56, 3.85]
ROWY = [1.66, 2.587, 3.513]
rect(GX, GY, GW, GH, WHITE, GRID, 1.0)
rect(COLX[1], GY + 0.008, 0.008, GH - 0.016, GRID, None, 0)
for ry in ROWY[1:]:
    rect(GX + 0.008, ry, GW - 0.016, 0.008, GRID, None, 0)

prods = [
    ("TestStand", "测试执行与管理",
     "序列化组织测试步骤，支持**并行与多工位**执行；结果入库与报表生成；可调用 **LabVIEW / C / Python** 模块。"),
    ("LabVIEW", "图形化开发环境",
     "数据流图形编程；具备**仪器连接与驱动调用**能力；内建界面控件；**测试库与工具包**缩短开发调试周期。"),
    ("TSM", "半导体测试扩展",
     "TestStand 之上增加**工位与引脚映射**、多工位并行；对接**探针台 / 分选机**，结果以 **STDF** 输出。"),
    ("PXI 平台", "模块化硬件",
     "机箱 + 控制器 + 模块组合成测量系统；**背板触发与时钟同步**；**换模块即换仪器构成**。"),
    ("NI-VISA / IVI", "仪器抽象层",
     "统一会话访问 **GPIB / LXI / USB / PXI** 仪器；**类驱动实现同类仪器互换**，上层调用不随硬件变更。"),
    ("InstrumentStudio", "仪器软面板",
     "以**台式仪器面板**操作模块化仪器；面板**可保存可复用**，便于调试与人工验证。"),
]
for i, (nm, cat, fe) in enumerate(prods):
    cx = COLX[i % 2]
    cy = ROWY[i // 2]
    pad = 0.14
    txt(cx + pad, cy + 0.11, 3.00, 0.18, nm, 7.4, True, INK)
    txt(cx + pad, cy + 0.295, 3.00, 0.15, cat, 5.7, False, MUT)
    txt_hi(cx + pad, cy + 0.465, 3.02, 0.40, fe, 6.2, BODY, BLUE2, True, 1.20)

# ── 右：产品拓扑
RX, RW = 7.44, 5.32
shead(RX, 1.36, 2.70, "产品拓扑：半导体产线联动")
txt(RX + 2.80, 1.39, RW - 2.80, 0.16, "图为 NI 公开产品图", 5.4, False, DIM, PP_ALIGN.RIGHT)

PIC_W, TX_W, PIC_H = 2.62, 2.54, 0.48
L1, L2, L3, L4, L5 = 1.66, 2.24, 2.82, 3.40, 3.96

picnode(RX, L1, PIC_W, PIC_H, "n2_ts.jpg", "TestStand", "执行与工位调度")
picnode(RX + PIC_W + 0.14, L1, TX_W, PIC_H, "n2_tsm.jpg", "TSM", "引脚映射 · 多工位")
wire([(RX + PIC_W, L1 + PIC_H / 2), (RX + PIC_W + 0.14, L1 + PIC_H / 2)])

picnode(RX, L2, PIC_W, PIC_H, "n2_lv.jpg", "LabVIEW", "测量程序 · 面板")
txtnode(RX + PIC_W + 0.14, L2, TX_W, PIC_H, "MES / STDF", "结果入库与产线回流")

SPINE = RX + PIC_W / 2
MID = RX + PIC_W + 0.14 + TX_W / 2
wire([(SPINE, L1 + PIC_H), (SPINE, L2)])
wire([(SPINE, L1 + PIC_H), (SPINE, L2 - 0.055), (MID, L2 - 0.055), (MID, L2)])

txtnode(RX, L3, RW, 0.44, "NI-VISA / IVI：统一会话与类驱动，屏蔽 GPIB / LXI / USB / PXI 差异")
wire([(SPINE, L2 + PIC_H), (SPINE, L3)])

picnode(RX, L4, PIC_W, PIC_H, "n2_pxi.jpg", "PXI 机箱与模块", "背板触发 · 时钟同步")
picnode(RX + PIC_W + 0.14, L4, TX_W, PIC_H, "n2_studio.jpg", "InstrumentStudio", "软面板调试与验证")
wire([(SPINE, L3 + 0.44), (SPINE, L4)])
wire([(SPINE, L3 + 0.44), (SPINE, L4 - 0.055), (MID, L4 - 0.055), (MID, L4)])

txtnode(RX, L5, 3.30, 0.48, "半导体测试台：探针台 / 分选机", "与 TSM 交换工位与分选信号")
txtnode(RX + 3.42, L5, 1.90, 0.48, "DUT：晶圆 / 芯片")
wire([(SPINE, L4 + PIC_H), (SPINE, L5)])
wire([(RX + 3.30, L5 + 0.24), (RX + 3.42, L5 + 0.24)])

# ── 核心思想（DSL 并入第 4 条）
shead(0.56, 4.56, 2.10, "核心思想")

ideas = [
    ("分布式拓扑", "控制、测量与数据服务分布在多台设备上，经工业以太网与总线按接口协作；工位可增减，单点故障影响面受限，扩容以增加节点实现。"),
    ("模块化与复用", "机箱、模块与背板同步构成复用基础；多种测量需求以模块组合覆盖，而非逐项定制整机。"),
    ("开放标准与生态", "PXI、VISA / IVI、STDF 等公开标准，使第三方模块、探针台与分选机均可接入，不被单一供应商锁定。"),
    ("描述与执行分离",
     "测试用**脚本（DSL）**声明意图——施加、测量、判定与流程，不写仪器细节；由执行引擎映射到通道，**换硬件不改脚本**，仿真与实机共用同一套描述。"),
]
bx, bw, gap = 0.56, 2.95, 0.133
for i, (ti, de) in enumerate(ideas):
    x = bx + i * (bw + gap)
    rect(x, 4.84, bw, 1.24, WHITE, GRID, 0.9)
    vrule(x, 4.84, 1.24, BLUE, 0.04)
    txt(x + 0.14, 4.91, 0.42, 0.22, "0" + str(i + 1), 10.5, True, RGBColor(0xC9, 0xC4, 0xB8))
    txt(x + 0.14, 5.16, bw - 0.28, 0.18, ti, 7.4, True, INK)
    txt_hi(x + 0.14, 5.39, bw - 0.28, 0.62, de, 6.1, BODY, BLUE2, True, 1.16)

# ── 小结
hline(0.56, 6.44, 12.20, RULE, 0.012)
txt(0.56, 6.50, 0.52, 0.20, "小结", 8.0, True, RED)
txt_hi(1.12, 6.49, 11.60, 0.44,
       "半导体场景下分工明确：**TestStand 编排工位与流程**、**TSM 承担引脚映射与多工位并行**、**驱动层对接探针台与分选机**、"
       "结果以 **STDF 回流 MES**；对应本方案，即 **DSL 描述与执行引擎分离**、仪器访问抽象、仿真与实机共用同一描述。",
       7.2, RGBColor(0x25, 0x25, 0x25), BLUE2, True, 1.18)

txt(0.56, 7.10, 11.6, 0.20,
    "资料来源：NI 公开公司与产品资料（TestStand、LabVIEW、TestStand Semiconductor Module、PXI、NI-VISA、InstrumentStudio）；"
    "产品特性为公开资料归纳，拓扑与核心思想为本文判断。", 6.0, False, DIM)
txt(12.30, 7.15, 0.50, 0.20, "1 / 1", 8, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
pics = oob = ov = 0
for si, s in enumerate(p.slides):
    bxs = []
    for sh in s.shapes:
        if sh.shape_type == 13:
            pics += 1
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bxs.append((l, t, w, h, sh.text_frame.text[:10]))
    for i in range(len(bxs)):
        for j in range(i + 1, len(bxs)):
            a, b = bxs[i], bxs[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜图片 {pics}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
