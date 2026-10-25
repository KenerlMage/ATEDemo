# -*- coding: utf-8 -*-
"""友商调研 · 制造测试装备篇（5 页，含自制图，允许贴图）"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "ASML制造测试装备.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x18, 0x18, 0x1A)
MUT = RGBColor(0x6E, 0x6C, 0x68)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x0F, 0x4C, 0x81)
GREEN = RGBColor(0x2F, 0x6B, 0x57)
AMBER = RGBColor(0xB0, 0x76, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD6, 0xD4, 0xCE)
SOFT = RGBColor(0xF8, 0xF8, 0xF6)
BAND = RGBColor(0xF7, 0xF4, 0xF0)
CREAM = RGBColor(0xFF, 0xF8, 0xEC)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _ea(run, size_pt, bold, color):
    f = run.font
    f.size = Pt(size_pt)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(sl, x, y, w, h, text, size=11, bold=False, color=INK,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        _ea(r, size, bold, color)
    return tb


def box(sl, x, y, w, h, fill, line=None, radius=0.10, lw=1.1):
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                              Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.text_frame.text = ""
    return shp


def arrow(sl, x, y, w, h, color, shape=MSO_SHAPE.RIGHT_ARROW):
    shp = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.text_frame.text = ""
    return shp


def header(sl, title, sub, page=None):
    box(sl, 0, 0, SW, 0.10, RED, None, 0.0)
    txt(sl, 0.55, 0.32, 11.3, 0.52, title, 24, True, INK)
    txt(sl, 0.57, 0.92, 11.7, 0.32, sub, 11.5, False, MUT)
    if page:
        txt(sl, 12.35, 7.12, 0.55, 0.24, f"{page}", 9, False, MUT, PP_ALIGN.RIGHT)


def footer(sl, s):
    txt(sl, 0.55, 7.10, 11.6, 0.26, s, 8, False, RGBColor(0x96, 0x94, 0x8E))


# ─────────────────────────── P1 重新定义问题
s1 = prs.slides.add_slide(BLANK)
header(s1, "造光刻机的测试装备：换个角度看 ASML",
       "ASML 不自己造所有零件，测试因此被分散在三个地方——不是一个车间、一台机器", 1)

stages = [
    ("① 模块设计与制造", "Carl Zeiss SMT / TRUMPF / VDL", "光学件、光源模块、精密结构件",
     "在供应商厂内完成零件级与模块级测试", RED),
    ("② 电子与电气集成", "Prodrive / Neways 等", "驱动板、控制板、线束、接口模块",
     "ICT、FCT、温循老化，平台化 ATE 的主场", AMBER),
    ("③ 总装与集成测试", "ASML Veldhoven 洁净室", "整机（光路、工件台、控制系统）",
     "集成、qualification、SPC 与 PFMEA", GREEN),
]
for i, (name, who, obj, test, col) in enumerate(stages):
    y = 1.52 + i * 1.62
    box(s1, 0.55, y, 6.40, 1.44, WHITE, LINE, 0.07, 1.0)
    box(s1, 0.55, y, 0.07, 1.44, col, None, 0.0)
    txt(s1, 0.74, y + 0.13, 5.9, 0.26, name, 12, True, col)
    txt(s1, 0.74, y + 0.44, 3.1, 0.22, who, 8.5, False, MUT)
    txt(s1, 0.74, y + 0.72, 5.9, 0.22, obj, 9, False, INK)
    txt(s1, 0.74, y + 1.00, 5.9, 0.30, test, 9, False, RGBColor(0x44, 0x42, 0x3E))
    if i < 2:
        arrow(s1, 3.55, y + 1.46, 0.22, 0.10, RGBColor(0xC0, 0xB8, 0xAC), MSO_SHAPE.DOWN_ARROW)

box(s1, 7.20, 1.52, 5.58, 2.44, CREAM, AMBER, 0.07, 1.4)
txt(s1, 7.42, 1.66, 5.15, 0.28, "本章的范围", 12, True, AMBER)
txt(s1, 7.42, 2.00, 5.15, 1.80,
    "只讨论「部件与模块怎么被自动测」：\n"
    "L1 光学件 · L2 电子模块 · L3 机电模块\n\n"
    "不讨论整机验收（L4）——那是设备商在客户厂房里做的事，与 ATE 不是同一类问题。",
    9.5, False, RGBColor(0x4A, 0x40, 0x30), spacing=1.25)

box(s1, 7.20, 4.14, 5.58, 2.62, BAND, RGBColor(0xE0, 0xDA, 0xD0), 0.07, 1.0)
txt(s1, 7.42, 4.28, 5.15, 0.26, "三条先记住的结论", 11.5, True, INK)
for i, t in enumerate([
    "层越高越像 ATE，层越低越像计量。",
    "同一句「自动测试」，四层由四种完全不同的装备承担。",
    "真正与我们同构的是 L2 与 L3：多品种、小批量、夹具各异。",
]):
    txt(s1, 7.42, 4.66 + i * 0.62, 5.15, 0.54, f"{i+1}. {t}", 9.5, False,
        RGBColor(0x3C, 0x3A, 0x36), spacing=1.2)

footer(s1, "依据：ZEISS SMT 资料、Prodrive / Neways 制造服务页、ASML 招聘岗位描述　｜　分层为本文判断，详见 05、06 两篇")

# ─────────────────────────── P2 四层制造链（贴图）
s2 = prs.slides.add_slide(BLANK)
header(s2, "制造链的四层，与每层在测什么",
       "模块化外包 + 洁净室总装：测试被拆到供应商厂、电子厂与总装厂三处", 2)
iw = 5.62 * 16 / 9
s2.shapes.add_picture(os.path.join(HERE, "ASML制造链与测试装备分层.png"),
                      Inches((SW - iw) / 2), Inches(1.22), width=Inches(iw))
txt(s2, 0.55, 6.88, 12.2, 0.26,
    "橙框标记的两层（L2 电子模块、L3 机电模块）是我们的目标场景；L1 是计量、L4 是系统验收，均不在 ATE 范畴。",
    9, False, MUT)
footer(s2, "图：ASML制造链与测试装备分层.png（自制，2400×1350）")

# ─────────────────────────── P3 装备谱系（贴图）
s3 = prs.slides.add_slide(BLANK)
header(s3, "四类测试手段 × 四层被测对象",
       "电性、计量、功能性能、环境可靠性——每类的装备与供应商都不同", 3)
s3.shapes.add_picture(os.path.join(HERE, "部件测试装备谱系矩阵.png"),
                      Inches((SW - iw) / 2), Inches(1.22), width=Inches(iw))
txt(s3, 0.55, 6.88, 12.2, 0.26,
    "橙框六格才是「装备制造用 ATE」的主战场；L1 的计量与 L4 的系统测试混进来，会把讨论带偏。",
    9, False, MUT)
footer(s3, "图：部件测试装备谱系矩阵.png（自制，2400×1350）　｜　具体测试项按通用工程实践归纳")

# ─────────────────────────── P4 三个判断
s4 = prs.slides.add_slide(BLANK)
header(s4, "三个判断：做哪一层、差什么、不做什么",
       "目标场景决定架构取舍，缺口决定投入顺序", 4)

txt(s4, 0.55, 1.42, 5.9, 0.26, "① 与量产 ATE 的分野", 11.5, True, BLUE)
rows = [("批量", "百万件级", "每月几十到几百件"),
        ("节拍", "压到秒级", "次要，允许分钟级"),
        ("被测件", "引脚定义固定", "每型号不同"),
        ("换型", "少但要极快", "频繁，可人工介入"),
        ("数据", "抽样统计", "全量可追溯")]
box(s4, 0.55, 1.72, 6.05, 0.34, RGBColor(0xEE, 0xEC, 0xE6), None, 0.06)
txt(s4, 0.68, 1.77, 1.2, 0.24, "维度", 8.5, True, MUT)
txt(s4, 1.95, 1.77, 2.15, 0.24, "量产 ATE", 8.5, True, MUT)
txt(s4, 4.20, 1.77, 2.30, 0.24, "装备制造用测试台", 8.5, True, MUT)
for i, (a, b, c) in enumerate(rows):
    ry = 2.10 + i * 0.52
    box(s4, 0.55, ry, 6.05, 0.46, WHITE if i % 2 == 0 else SOFT, LINE, 0.06, 0.8)
    txt(s4, 0.68, ry + 0.11, 1.2, 0.24, a, 9, True, INK)
    txt(s4, 1.95, ry + 0.11, 2.15, 0.24, b, 8.5, False, MUT)
    txt(s4, 4.20, ry + 0.11, 2.30, 0.24, c, 8.5, False, RGBColor(0x2C, 0x40, 0x34))
txt(s4, 0.55, 4.80, 6.05, 0.44,
    "结论：我们的架构匹配度在右侧那一列——不是性能不够，而是场景本来就不同。", 9, False, MUT)

txt(s4, 6.95, 1.42, 5.8, 0.26, "② 现有形态的四个缺口", 11.5, True, AMBER)
gaps = [("site / 工位抽象", "无，台位靠人工约定 → 一拖多无法并行"),
        ("夹具与接线描述", "只有仪器地址，没有「测什么接哪个针」"),
        ("非电信号数据模型", "波形、图像、面形只能存路径，无法比较"),
        ("参数闭环", "只判 pass/fail，偏差变不成下一轮参数")]
for i, (t, s) in enumerate(gaps):
    gy = 1.72 + i * 0.86
    box(s4, 6.95, gy, 5.80, 0.74, WHITE, LINE, 0.07, 0.9)
    box(s4, 6.95, gy, 0.06, 0.74, AMBER, None, 0.0)
    txt(s4, 7.12, gy + 0.08, 5.45, 0.24, t, 9.5, True, INK)
    txt(s4, 7.12, gy + 0.38, 5.45, 0.26, s, 8.5, False, MUT)

txt(s4, 6.95, 5.30, 5.80, 0.26, "③ 明确不做", 11.5, True, RED)
box(s4, 6.95, 5.60, 5.80, 1.14, RGBColor(0xFA, 0xF6, 0xF4), RGBColor(0xE2, 0xD4, 0xD0), 0.07, 1.0)
txt(s4, 7.12, 5.70, 5.46, 0.94,
    "不做 L1 的计量仪器（干涉仪是另一个行业）；\n"
    "不做 L4 的整机验收（边界在客户厂房）；\n"
    "不与专用 per-pin ATE 比性能（百 MHz 级 PXI 已够）。",
    9, False, RGBColor(0x5A, 0x46, 0x42), spacing=1.3)
footer(s4, "判断依据：ASML 制造链分层（05）与测试装备谱系（06）；数值为工程通例，非厂商内部数据")

# ─────────────────────────── P5 来源与口径
s5 = prs.slides.add_slide(BLANK)
header(s5, "来源与口径：哪些是事实，哪些是判断",
       "这一页建议保留——它是这份材料能否被追问的基础", 5)

box(s5, 0.55, 1.46, 6.05, 4.72, WHITE, LINE, 0.06, 1.0)
txt(s5, 0.75, 1.60, 5.6, 0.26, "公开事实（可直接引用）", 11.5, True, BLUE)
facts = [
    ("Zeiss 镜片工序链", "抛光 → CCOS → 离子束修形 → 镀膜 → 面形量测"),
    ("Zeiss 计量要求", "面形误差与粗糙度同时达标，靠干涉仪等计量设备支撑"),
    ("Prodrive 制造服务", "列出全流程测试自动化、温湿度测试、洁净室装配"),
    ("Neways 半导体能力", "6,700 m² 以上洁净室（最高 ISO 5）、微电子封装"),
    ("ASML 集成测试岗位", "集成多子系统、qualification、SPC 与 PFMEA"),
    ("供应链结构", "光学 Zeiss、激光 TRUMPF、精密加工 VDL 等一级供应商"),
]
for i, (t, s) in enumerate(facts):
    y = 1.94 + i * 0.70
    txt(s5, 0.75, y, 5.6, 0.24, t, 9.5, True, INK)
    txt(s5, 0.75, y + 0.26, 5.6, 0.36, s, 8.5, False, MUT)

box(s5, 6.80, 1.46, 5.98, 4.72, CREAM, AMBER, 0.06, 1.2)
txt(s5, 7.00, 1.60, 5.5, 0.26, "本文判断（须标明）", 11.5, True, AMBER)
judges = [
    ("四层划分", "L1 光学件 / L2 电子模块 / L3 机电模块 / L4 整机"),
    ("装备归属映射", "哪一层由 ATE、哪一层由计量、哪一层由系统测试承担"),
    ("测试项清单", "按半导体设备制造通用工程实践归纳"),
    ("我们的切入点", "L2 / L3 的电性、功能性能与环境可靠性测试"),
]
for i, (t, s) in enumerate(judges):
    y = 1.94 + i * 0.96
    txt(s5, 7.00, y, 5.5, 0.24, t, 9.5, True, INK)
    txt(s5, 7.00, y + 0.26, 5.5, 0.60, s, 8.5, False, RGBColor(0x4A, 0x40, 0x30))

box(s5, 0.55, 6.32, 12.23, 0.56, BAND, RGBColor(0xE0, 0xDA, 0xD0), 0.08, 1.0)
txt(s5, 0.75, 6.44, 11.9, 0.34,
    "汇报口径：讲「这类测试在工程上如何分层、每层需要什么能力」，不讲「ASML 某台设备的具体型号」——后者无公开资料。",
    9, False, RGBColor(0x4A, 0x46, 0x40))
footer(s5, "完整引用清单见 docs/友商调研/README.md（R1–R16）；图片来源为自制图表，非网页截图")

prs.save(OUT)

# ── 自检
p = Presentation(OUT)
pics = oob = ov = 0
for si, sl in enumerate(p.slides):
    boxes = []
    for sh in sl.shapes:
        if sh.shape_type == 13:
            pics += 1
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界 p{si+1}: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            boxes.append((l, t, w, h, sh.text_frame.text[:14]))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.06 and oy > 0.06:
                ov += 1
                print(f"  重叠 p{si+1}: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜图片 {pics}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
