# -*- coding: utf-8 -*-
"""文本容量自检：按字宽估算每行字数与所需高度，找出溢出文本框"""
import math
import sys
from pptx import Presentation
from pptx.util import Pt


def wide(ch):
    if ord(ch) < 128:
        return 0.53
    return 1.0


p = Presentation(sys.argv[1])
bad = 0
for si, sl in enumerate(p.slides):
    for sh in sl.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        wpt = sh.width.inches * 72
        hpt = sh.height.inches * 72
        need = 0.0
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs)
            if not t:
                continue
            size = None
            for r in para.runs:
                if r.font.size:
                    size = r.font.size.pt
                    break
            size = size or 12
            per = max(1.0, wpt / size)
            units = sum(wide(c) for c in t)
            lines = max(1, math.ceil(units / per))
            ls = para.line_spacing or 1.18
            need += lines * size * ls
        if need > hpt + 1.0:
            bad += 1
            print(f" p{si+1} 溢出 {need:.0f}pt > {hpt:.0f}pt 「{sh.text_frame.text[:26]}」")
print(f"文本溢出 {bad} 处")
