# -*- coding: utf-8 -*-
"""业界方案洞察 · 一页归并版：公司 → 产品与产品特性 → 核心思想"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "业界方案洞察-NI-一页归并版.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
RULE = RGBColor(0xBF, 0xBB, 0xB1)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR = RGBColor(0xEE, 0xEC, 0xE6)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14, italic=False):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = ln
        _font(r, size, bold, color, italic)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.012):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(x, y, h, color=RED, w=0.045):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def secno(x, y, no, name):
    txt(x, y - 0.04, 0.34, 0.30, no, 13.5, True, RED)
    txt(x + 0.40, y + 0.02, 1.70, 0.18, name, 8.0, True, RED)
    hline(x + 0.40, y + 0.24, 1.66, RULE, 0.010)


# ── 页头
txt(0.56, 0.22, 7.80, 0.38, "业界方案洞察", 21, True, INK)
txt(8.60, 0.34, 4.16, 0.22, "NI：公司 · 产品与特性 · 核心思想", 8.4, False, DIM, PP_ALIGN.RIGHT)

# ── 一、公司
rect(0.56, 0.62, 12.20, 0.86, GRAY, None, 0)
vrule(0.56, 0.62, 0.86, RED, 0.05)
txt(0.72, 0.70, 0.30, 0.30, "一", 13.5, True, RED)
txt(0.72, 1.06, 0.60, 0.18, "公司", 7.2, True, RED)
txt(1.42, 0.70, 11.20, 0.72,
    "美国国家仪器（National Instruments，NI）1976 年成立于美国得克萨斯州奥斯汀，2023 年并入 Emerson，作为其测试与测量业务运营。"
    "公司以测量与控制软件平台为核心，配合模块化硬件，服务半导体、汽车、航空航天与国防、能源、通信及院校科研等领域的验证与生产测试。"
    "1986 年发布 LabVIEW，由此确立「虚拟仪器」路线：以软件定义仪器功能，硬件只提供通用的采集与激励能力。",
    7.2, False, RGBColor(0x2B, 0x2B, 0x2B), ls=1.30)

# ── 二、产品与产品特性
secno(0.56, 1.62, "二", "产品与产品特性")

rect(0.56, 1.94, 12.20, 0.24, HDR, None, 0)
txt(0.68, 1.975, 1.55, 0.16, "产品", 6.5, True, MUT)
txt(2.24, 1.975, 1.75, 0.16, "类别与定位", 6.5, True, MUT)
txt(4.10, 1.975, 8.50, 0.16, "关键特性（官方口径与公开资料）", 6.5, True, MUT)

rows = [
    ("TestStand", "测试执行与管理软件",
     "以序列组织测试步骤；支持并行与多工位执行；记录结果并生成报表；可调用 LabVIEW、C/C#、Python 等代码模块；不绑定具体仪器。"),
    ("LabVIEW", "图形化开发环境",
     "数据流图形编程；具备仪器连接与驱动调用能力；内建用户界面控件；配测试专用库与工具包，缩短开发与调试周期。"),
    ("TSM", "半导体测试扩展",
     "在 TestStand 之上增加半导体测试所需的工位与引脚映射、多工位并行及量产流程支持；官方定位为开发、调试、部署与维护半导体测试系统。"),
    ("PXI 平台", "模块化硬件",
     "机箱、控制器与模块组合成测量系统；背板提供触发与时钟同步；按测量类型更换模块即可改变仪器构成，无需专用整机。"),
    ("NI-VISA / IVI", "仪器抽象层",
     "以统一会话访问 GPIB、LXI、USB、串口与 PXI 仪器；IVI 提供类驱动，实现同类仪器互换，上层调用不随硬件变更。"),
    ("InstrumentStudio", "仪器软面板",
     "以台式仪器面板方式操作模块化仪器；面板可保存与复用，便于调试与人工验证，减少自编界面工作量。"),
]
ry = 2.18
for i, (nm, cat, fe) in enumerate(rows):
    if i % 2 == 0:
        rect(0.56, ry, 12.20, 0.40, WHITE, None, 0)
    else:
        rect(0.56, ry, 12.20, 0.40, RGBColor(0xFA, 0xF9, 0xF6), None, 0)
    hline(0.56, ry, 12.20, RGBColor(0xDD, 0xD9, 0xD0), 0.008)
    txt(0.68, ry + 0.115, 1.56, 0.18, nm, 7.3, True, INK)
    txt(2.24, ry + 0.115, 1.80, 0.18, cat, 6.4, False, BLUE)
    txt(4.10, ry + 0.055, 8.50, 0.32, fe, 6.3, False, BODY, ls=1.12)
    ry += 0.40
hline(0.56, ry, 12.20, RGBColor(0xDD, 0xD9, 0xD0), 0.010)

# ── 三、核心思想
secno(0.56, 4.76, "三", "核心思想")

ideas = [
    ("软件定义仪器", "功能由软件配置决定，硬件只提供通用采集与激励；同一机箱加载不同程序即成为不同仪器。",
     "启示：把测量能力放进软件层，硬件趋于通用。"),
    ("模块化与复用", "机箱、模块与背板同步构成复用基础；多种测量需求以组合覆盖，而非逐项定制整机。",
     "启示：以组合覆盖品类，降低小批量成本。"),
    ("开放标准与生态", "PXI、VISA/IVI、ATML 等公开标准，使第三方模块与驱动可接入，用户不被单一供应商锁定。",
     "启示：接口按标准实现，保留替换空间。"),
    ("描述与执行分离", "TestStand 负责编排，LabVIEW 与代码模块负责执行，驱动负责适配仪器，三者互不越界。",
     "启示：测试描述与执行引擎解耦，换平台不改描述。"),
]
bx, bw, gap = 0.56, 2.95, 0.133
for i, (ti, de, hint) in enumerate(ideas):
    x = bx + i * (bw + gap)
    rect(x, 5.08, bw, 1.24, WHITE, RULE, 0.75)
    vrule(x, 5.08, 1.24, BLUE, 0.04)
    txt(x + 0.13, 5.15, 0.42, 0.26, "0" + str(i + 1), 11.5, True, RGBColor(0xC9, 0xC4, 0xB8))
    txt(x + 0.13, 5.42, bw - 0.26, 0.18, ti, 7.6, True, INK)
    txt(x + 0.13, 5.63, bw - 0.26, 0.42, de, 6.2, False, BODY, ls=1.12)
    hline(x + 0.13, 6.06, bw - 0.26, RGBColor(0xDD, 0xD9, 0xD0), 0.008)
    txt(x + 0.13, 6.10, bw - 0.26, 0.18, hint, 6.0, False, BLUE)

# ── 小结
hline(0.56, 6.44, 12.20, RULE, 0.010)
txt(0.56, 6.50, 0.52, 0.20, "小结", 8.0, True, RED)
txt(1.12, 6.50, 11.60, 0.44,
    "NI 的可借鉴之处在于分层与解耦：以软件平台为价值核心，以模块化硬件承载测量能力，以公开标准维持生态；"
    "对应到本方案，即测试描述与执行引擎分离、仪器访问抽象为统一驱动、仿真与实机共用同一套描述。",
    7.6, False, RGBColor(0x25, 0x25, 0x25), ls=1.2)

txt(0.56, 7.09, 11.4, 0.20,
    "资料来源：NI 公开公司信息与产品页（TestStand、LabVIEW、TestStand Semiconductor Module、PXI、NI-VISA、InstrumentStudio）；"
    "产品特性为公开资料归纳，思想归纳与启示为本文判断。", 6.2, False, DIM)
txt(12.30, 7.14, 0.50, 0.20, "1 / 1", 8, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
oob = ov = 0
for si, s in enumerate(p.slides):
    bx2 = []
    for sh in s.shapes:
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bx2.append((l, t, w, h, sh.text_frame.text[:12]))
    for i in range(len(bx2)):
        for j in range(i + 1, len(bx2)):
            a, b = bx2[i], bx2[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
