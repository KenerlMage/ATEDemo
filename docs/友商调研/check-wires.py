# -*- coding: utf-8 -*-
"""连线连通性审计：端点是否落在节点上、线段是否穿过无关节点"""
import sys
from pptx import Presentation
from pptx.util import Emu

INSET = 0.04


def rects(sl):
    out = []
    for sh in sl.shapes:
        tag = sh._element.tag.split("}")[-1]
        if tag != "sp":
            continue
        try:
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if w < 1.2 or h < 0.40:
            continue
        out.append((sh.left.inches, sh.top.inches, w, h,
                    (sh.text_frame.text[:14] if sh.has_text_frame else "")))
    return out


def conns(sl):
    out = []
    for sh in sl.shapes:
        if sh._element.tag.split("}")[-1] != "cxnSp":
            continue
        xfrm = sh._element.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
        spPr = sh._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
        if spPr is None:
            continue
        x = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
        if x is None:
            continue
        off = x.find("{http://schemas.openxmlformats.org/drawingml/2006/main}off")
        ext = x.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ext")
        cx = int(off.get("x")) / 914400
        cy = int(off.get("y")) / 914400
        cw = int(ext.get("cx")) / 914400
        ch = int(ext.get("cy")) / 914400
        fl = x.get("flipH") == "1"
        fv = x.get("flipV") == "1"
        x1, y1 = (cx + cw, cy) if fl else (cx, cy)
        x2, y2 = (cx, cy + ch) if fv else (cx + cw, cy + ch)
        if fl and fv:
            x1, y1, x2, y2 = cx + cw, cy + ch, cx, cy
        out.append((x1, y1, x2, y2))
    return out


def on_edge(px, py, r):
    x, y, w, h, _ = r
    if not (x - 0.06 <= px <= x + w + 0.06 and y - 0.06 <= py <= y + h + 0.06):
        return False
    return (abs(px - x) <= 0.06 or abs(px - (x + w)) <= 0.06 or
            abs(py - y) <= 0.06 or abs(py - (y + h)) <= 0.06)


def crosses(seg, r):
    x1, y1, x2, y2 = seg
    x, y, w, h, _ = r
    x0, y0 = x + INSET, y + INSET
    x3, y3 = x + w - INSET, y + h - INSET
    if x3 <= x0 or y3 <= y0:
        return False
    if abs(y1 - y2) < 1e-6:                      # 水平
        lo, hi = sorted((x1, x2))
        return y0 < y1 < y3 and lo < x3 and hi > x0
    if abs(x1 - x2) < 1e-6:                      # 垂直
        lo, hi = sorted((y1, y2))
        return x0 < x1 < x3 and lo < y3 and hi > y0
    return False


p = Presentation(sys.argv[1])
issues = 0
for si, sl in enumerate(p.slides):
    rs = rects(sl)
    for seg in conns(sl):
        x1, y1, x2, y2 = seg
        # 端点是否贴着某个节点
        for px, py in ((x1, y1), (x2, y2)):
            if not any(on_edge(px, py, r) for r in rs):
                print(f" p{si+1} 端点悬空 ({px:.2f},{py:.2f})")
                issues += 1
        # 线段中段是否穿过别的节点
        for r in rs:
            if crosses(seg, r):
                print(f" p{si+1} 线穿节点 ({x1:.2f},{y1:.2f})→({x2:.2f},{y2:.2f}) "
                      f"× 「{r[4]}」")
                issues += 1
print(f"连线问题 {issues} 处")
