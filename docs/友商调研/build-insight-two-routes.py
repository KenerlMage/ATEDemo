# -*- coding: utf-8 -*-
"""业界方案洞察 · v5：一页两家公司、两种思路（左 NI 平台化 / 右 南方公司 中心化下发）"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "业界方案洞察-两种思路.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
BLUE2 = RGBColor(0x1B, 0x5E, 0xA8)
TEAL = RGBColor(0x0F, 0x5E, 0x5A)
RULE = RGBColor(0xAD, 0xA8, 0x9C)
GRID = RGBColor(0xC6, 0xC1, 0xB5)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR = RGBColor(0xEE, 0xEC, 0xE6)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def para(tf, first, segs, size, bold, color, align, ls):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = ls
    for t, c, b in segs:
        r = p.add_run()
        r.text = t
        _font(r, size, bold if b is None else b, color if c is None else c)
    return p


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        para(tf, i == 0, [(ln, None, None)], size, bold, color, align, ls)
    return tb


def txt_hi(x, y, w, h, text, size=6.2, color=BODY, hi=BLUE2, hi_bold=True,
           ls=1.18, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    segs = []
    for i, seg in enumerate(str(text).split("**")):
        if not seg:
            continue
        on = (i % 2 == 1)
        segs.append((seg, hi if on else None, hi_bold if on else None))
    para(tf, True, segs, size, False, color, align, ls)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.010):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(x, y, h, color=RED, w=0.040):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vline(x, y, h, color=RULE, w=0.008):
    rect(x, y, w, h, color, None, 0)


def shead(x, y, w, text, color=RED, size=7.0):
    txt(x, y, w, 0.18, text, size, True, color)
    hline(x, y + 0.19, w, GRID, 0.010)


# ── 页头
txt(0.56, 0.16, 8.40, 0.34, "业界方案洞察：两种思路", 19, True, INK)
txt(9.20, 0.24, 3.56, 0.22, "NI 平台化 · 南方公司 中心化下发", 8.0, False, DIM, PP_ALIGN.RIGHT)
hline(0.56, 0.54, 12.20, RULE, 0.014)

LX, LW = 0.56, 5.74
RX2, RW2 = 6.86, 5.90
TOPY = 0.64
vline(6.58, 0.66, 6.02, GRID, 0.008)

# ══════════ 左半：NI ══════════
txt(LX, TOPY, 3.20, 0.22, "NI　美国国家仪器", 9.6, True, INK)
txt(LX + 3.30, TOPY + 0.03, LW - 3.30, 0.18, "平台化路线", 6.6, False, DIM, PP_ALIGN.RIGHT)

rect(LX, 0.90, LW, 0.60, GRAY, None, 0)
vrule(LX, 0.90, 0.60, RED, 0.045)
txt_hi(LX + 0.16, 0.975, LW - 0.30, 0.48,
       "1976 年成立于奥斯汀，2023 年并入 Emerson；以**软件平台**为核心、**模块化硬件**为承载，"
       "服务半导体、汽车、航空航天与院校科研；1986 年发布 LabVIEW，确立**虚拟仪器**路线。",
       6.5, RGBColor(0x2B, 0x2B, 0x2B), BLUE2, True, 1.24)

shead(LX, 1.58, LW, "产品与产品特性")
prods = [
    ("TestStand", "序列编排、**并行与多工位**、结果入库"),
    ("LabVIEW", "**图形化开发**；仪器连接与驱动调用"),
    ("TSM", "**工位与引脚映射**；对接探针台、**STDF**"),
    ("PXI 平台", "机箱加模块，**背板同步**，换模块换仪器"),
    ("NI-VISA / IVI", "**统一会话** GPIB/LXI/PXI，**类驱动互换**"),
    ("InstrumentStudio", "**台式面板**操作模块，面板可复用"),
]
CELLW, CELLH = 2.80, 0.72
for i, (nm, fe) in enumerate(prods):
    cx = LX + (i % 2) * (CELLW + 0.14)
    cy = 1.82 + (i // 2) * 0.76
    rect(cx, cy, CELLW, CELLH, WHITE, GRID, 0.7)
    txt(cx + 0.11, cy + 0.075, CELLW - 0.18, 0.16, nm, 6.6, True, INK)
    txt_hi(cx + 0.11, cy + 0.245, CELLW - 0.18, 0.40, fe, 5.7, BODY, BLUE2, True, 1.14)

shead(LX, 4.14, LW, "核心思想")
ideas = [
    ("01", "分布式拓扑", "设备分散，按接口协作"),
    ("02", "模块化复用", "组合覆盖品类需求"),
    ("03", "开放标准", "PXI、VISA/IVI、STDF"),
    ("04", "描述与执行分离", "脚本（DSL）声明意图"),
]
for i, (no, ti, de) in enumerate(ideas):
    cx = LX + (i % 2) * (CELLW + 0.14)
    cy = 4.40 + (i // 2) * 0.44
    txt(cx, cy, 0.26, 0.16, no, 6.0, True, RGBColor(0xC9, 0xC4, 0xB8))
    txt(cx + 0.28, cy, CELLW - 0.28, 0.16, ti, 6.4, True, INK)
    txt(cx + 0.28, cy + 0.185, CELLW - 0.28, 0.16, de, 5.7, False, MUT)

# ══════════ 右半：南方公司 ══════════
txt(RX2, TOPY, 3.24, 0.22, "南方公司（华为）", 9.6, True, INK)
txt(RX2 + 3.30, TOPY + 0.03, RW2 - 3.30, 0.18, "中心化下发路线", 6.6, False, DIM, PP_ALIGN.RIGHT)

rect(RX2, 0.90, RW2, 0.60, GRAY, None, 0)
vrule(RX2, 0.90, 0.60, TEAL, 0.045)
txt_hi(RX2 + 0.16, 0.975, RW2 - 0.30, 0.48,
       "面向**终端设备制造**与**通信**领域；测试软件集中在中心侧维护，按工位与槽位实时下发，"
       "以**统一基线**保证多地产线运行同一版本，测试动作以脚本编排。",
       6.5, RGBColor(0x2B, 0x2B, 0x2B), TEAL, True, 1.24)

shead(RX2, 1.58, RW2, "方案要点", color=BLUE)
points = [
    ("DSL 驱动设备", "以脚本描述测试动作，执行端解释并驱动仪器与柜体设备。"),
    ("中心化部署", "测试软件集中于中心，版本统一维护，不散落在各工位。"),
    ("实时下发机柜—框—槽位", "按工位与槽位粒度推送软件与配置，改版即下发。"),
    ("脚本高度复用", "同一脚本跨机型、跨工位复用，减少重复编写。"),
    ("动作抽象与编排", "多数测试动作抽象为原语，按流程自由编排组合。"),
]
for i, (ti, de) in enumerate(points):
    cy = 1.82 + i * 0.448
    txt(RX2, cy, 0.24, 0.16, str(i + 1), 6.2, True, TEAL)
    txt(RX2 + 0.24, cy, 5.50, 0.16, ti, 6.6, True, INK)
    txt(RX2 + 0.24, cy + 0.185, 5.50, 0.20, de, 5.8, False, MUT)

shead(RX2, 4.14, RW2, "公司特性", color=BLUE)
feats = [
    ("软件基线程度高", "测试程序以基线版本统一管理，各工位运行同一基线，版本可追溯。"),
    ("执行自动化水平高", "从下发、执行到结果回收由系统完成，人工介入少、节拍稳定。"),
]
for i, (ti, de) in enumerate(feats):
    cx = RX2 + i * 2.96
    rect(cx, 4.40, 2.82, 0.78, FILL2, GRID, 0.8)
    vrule(cx, 4.40, 0.78, TEAL, 0.04)
    txt(cx + 0.12, 4.47, 2.60, 0.16, ti, 6.8, True, TEAL)
    txt(cx + 0.12, 4.66, 2.60, 0.46, de, 5.8, False, BODY, ls=1.16)

# ══════════ 下部：两种思路对照 ══════════
CY = 5.66
txt(0.56, CY, 3.00, 0.18, "两种思路对照", 7.2, True, RED)
hline(0.56, CY + 0.20, 12.20, GRID, 0.010)

rect(0.56, 5.94, 12.20, 0.22, HDR, None, 0)
txt(0.66, 5.975, 1.30, 0.16, "维度", 5.9, True, MUT)
txt(2.00, 5.975, 4.90, 0.16, "NI：平台化", 5.9, True, MUT)
txt(7.06, 5.975, 5.60, 0.16, "南方公司：中心化下发", 5.9, True, MUT)

rows = [
    ("部署形态", "分散：Host、机箱、仪器与测试台经网络协作", "集中：软件在中心，按槽位下发到机柜"),
    ("测试描述", "图形化程序与序列编排，驱动适配仪器", "DSL 脚本声明动作，动作抽象为原语"),
    ("复用方式", "模块组合复用硬件，标准接口换生态", "脚本与基线复用，中心改版整体下发"),
]
ry = 6.16
for a, b, c in rows:
    txt(0.66, ry + 0.035, 1.30, 0.16, a, 6.0, True, INK)
    txt(2.00, ry + 0.035, 4.86, 0.16, b, 5.9, False, BODY)
    txt(7.06, ry + 0.035, 5.60, 0.16, c, 5.9, False, BODY)
    ry += 0.22
    hline(0.56, ry - 0.005, 12.20, RGBColor(0xDD, 0xD9, 0xD0), 0.008)

txt(0.56, 6.88, 12.20, 0.20,
    "判断：差别不在工具，而在重心——NI 把能力放在可组合的硬件与开放接口上，适配多品种小批量；"
    "南方公司把能力放在可下发的软件基线与脚本复用上，适配机型集中、版本一致性要求高的产线。",
    6.2, False, RGBColor(0x3A, 0x38, 0x34))

txt(0.56, 7.14, 11.0, 0.20,
    "资料来源：NI 部分引自公开产品资料；南方公司部分为访谈与内部归纳，未引用其产品名称；路线判断为本文观点。",
    5.8, False, DIM)
txt(12.30, 7.16, 0.50, 0.20, "1 / 1", 7.6, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
oob = ov = 0
for si, s in enumerate(p.slides):
    bxs = []
    for sh in s.shapes:
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bxs.append((l, t, w, h, sh.text_frame.text[:10]))
    for i in range(len(bxs)):
        for j in range(i + 1, len(bxs)):
            a, b = bxs[i], bxs[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
