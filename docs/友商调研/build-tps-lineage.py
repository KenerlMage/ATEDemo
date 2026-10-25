# -*- coding: utf-8 -*-
"""在图标版基础上追加一页：TPS 脚本 DSL 的形制来源（ATLAS / IEEE 416·716 / 1641 / ATML）"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "业界方案洞察-NI-图标版.pptx")
OUT = os.path.join(HERE, "业界方案洞察-NI-含TPS谱系.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
RULE = RGBColor(0xBF, 0xBB, 0xB1)
FILL = RGBColor(0xF7, 0xF6, 0xF3)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WIRE = RGBColor(0x7A, 0x76, 0x6E)

prs = Presentation(SRC)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14, italic=False):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = ln
        _font(r, size, bold, color, italic)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.012):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(x, y, h, color=RED, w=0.045):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def wire(pts, color=WIRE, lw=1.0, arrow=True, dash=False):
    for i in range(len(pts) - 1):
        x1, y1 = pts[i]
        x2, y2 = pts[i + 1]
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(lw)
        ln = c.line._get_or_add_ln()
        if dash:
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "dash")
        if arrow and i == len(pts) - 2:
            t = etree.SubElement(ln, qn("a:tailEnd"))
            t.set("type", "triangle")
            t.set("w", "med")
            t.set("len", "med")


def snode(x, y, w, h, year, title, desc, accent=BLUE, fill=WHITE):
    rect(x, y, w, h, fill, accent, 0.9)
    vrule(x, y, h, accent, 0.045)
    txt(x + 0.11, y + 0.07, 1.05, 0.16, year, 5.9, True, DIM)
    txt(x + 0.11, y + 0.24, w - 0.22, 0.18, title, 7.2, True, INK)
    txt(x + 0.11, y + 0.43, w - 0.22, h - 0.47, desc, 5.9, False, MUT, ls=1.1)


def shead(x, y, w, text):
    txt(x, y, w, 0.17, text, 7.9, True, RED)
    hline(x, y + 0.21, w, RULE)


# ── 页头（与第一页同一版式）
txt(0.56, 0.22, 8.0, 0.38, "业界方案洞察", 21, True, INK)
txt(8.90, 0.34, 3.86, 0.22, "TPS 脚本的 DSL 形制及其来源", 8.4, False, DIM, PP_ALIGN.RIGHT)
rect(0.56, 0.66, 12.20, 0.70, GRAY, None, 0)
vrule(0.56, 0.66, 0.70, RED, 0.05)
txt(0.74, 0.75, 11.85, 0.54,
    "以类 DSL 描述测试步骤的 TPS，在业界有明确谱系：ATLAS 由航空无线电公司（ARINC）于 1960 年代为机载设备测试开发，"
    "后由 IEEE 标准化为 416 与 716（C/ATLAS）；2000 年后由 IEEE 1641（信号与测试定义）与 IEEE 1671（ATML）承接。"
    "半导体 ATE 侧未形成公开统一语言，厂商各自私有，如 Teradyne IG-XL 与 Advantest SmarTest。",
    7.3, False, RGBColor(0x2E, 0x2E, 0x2E), ls=1.26)

# ── 左：谱系链
CX, CW = 0.56, 3.34
chain = [
    ("1960s", "ARINC ATLAS", "航空无线电公司为机载设备测试开发，以信号与动作描述测试过程。", RED, FILL2),
    ("1984", "IEEE 416 ATLAS", "成为跨系统的通用测试语言标准，按级别定义语言子集。", BLUE, WHITE),
    ("1985", "IEEE 716 C/ATLAS", "面向商用测试系统的子集，用于 TPS 编写与移植。", BLUE, FILL2),
    ("2004", "IEEE 1641 STD", "以基本信号模型描述测试需求，替代 ATLAS 的表述方式。", BLUE, WHITE),
    ("2006", "IEEE 1671 ATML", "以 XML 交换测试信息：测试描述、仪器与结果。", BLUE, FILL2),
]
cy = 1.46
for yr, ti, de, ac, fi in chain:
    snode(CX, cy, CW, 0.78, yr, ti, de, ac, fi)
    if cy < 5.0:
        wire([(CX + CW / 2, cy + 0.78), (CX + CW / 2, cy + 0.94)])
    cy += 0.94

# ── 中部偏右：厂商私有形制（旁支）
BX, BW = 4.18, 3.10
snode(BX, 2.40, BW, 0.78, "同期", "厂商私有形制", "半导体与板级 ATE 未统一语言：\nTeradyne IG-XL 以图形化环境组织测试程序；\nAdvantest SmarTest 为 V93000 系统软件。", RED, FILL2)
wire([(CX + CW, 2.79), (BX, 2.79)], dash=True, arrow=False)

# ── 中部偏右：三点启示
tips = [
    ("① 语言与执行分离", "ATLAS 只描述信号与动作，由编译器映射到具体仪器；我们以 JSON 描述、pytest 执行，属同一分工。"),
    ("② 动词可对齐", "apply、measure、compare、connect 是 ATLAS 的核心动词，与 TPS 中的动作名一一对应。"),
    ("③ 有标准可引", "对外可称参照 IEEE 416 / 716 与 IEEE 1641 的表述方式，较「自研 DSL」更易被评审接受。"),
]
ty = 3.34
for ti, de in tips:
    txt(BX, ty, BW, 0.17, ti, 6.8, True, INK)
    txt(BX, ty + 0.18, BW, 0.44, de, 6.0, False, BODY, ls=1.12)
    ty += 0.86

# ── 右栏：术语对照
RX, RW = 7.52, 5.24
shead(RX, 1.44, RW, "术语对照：ATLAS 动词与 TPS 字段")
pairs = [
    ("APPLY / SETUP", "apply / setup", "施加激励、设置量程与通道"),
    ("MEASURE (MEAS)", "measure", "采集响应并换算为工程量"),
    ("COMPARE", "limit / compare", "与上下限比较，给出合格判定"),
    ("CONNECT / DISCONNECT", "connect / route", "接通或断开指定通路"),
]
rect(RX, 1.70, RW, 0.24, RGBColor(0xEE, 0xEC, 0xE6), None, 0)
txt(RX + 0.10, 1.735, 1.70, 0.16, "ATLAS 语句", 6.4, True, MUT)
txt(RX + 1.84, 1.735, 1.30, 0.16, "TPS 字段", 6.4, True, MUT)
txt(RX + 3.18, 1.735, RW - 3.24, 0.16, "含义", 6.4, True, MUT)
ry = 1.94
for a, b, c in pairs:
    rect(RX, ry, RW, 0.36, WHITE, RULE, 0.5)
    txt(RX + 0.10, ry + 0.09, 1.70, 0.18, a, 6.3, True, INK)
    txt(RX + 1.84, ry + 0.09, 1.30, 0.18, b, 6.3, False, BLUE)
    txt(RX + 3.18, ry + 0.09, RW - 3.24, 0.18, c, 6.2, False, MUT)
    ry += 0.36

# ── 右栏：出处
shead(RX, 3.48, RW, "可引用出处")
refs = [
    ("R1", "IEEE Std 416-1984（ATLAS）　ATLAS 最初为航空应用开发，由 ARINC 提出。"),
    ("R2", "IEEE Std 716-1985 / 1995（C/ATLAS）　416 的子集，面向商用自动测试系统。"),
    ("R3", "IEEE Std 1641-2004 / 2022（Signal and Test Definition）　承接 ATLAS 的信号描述。"),
    ("R4", "IEEE Std 1671（ATML）　以 XML 交换测试描述、仪器与结果信息。"),
    ("R5", "Teradyne IG-XL 产品页；Advantest SmarTest（V93000 系统软件）。"),
]
ry2 = 3.70
for k, t in refs:
    txt(RX, ry2, 0.34, 0.15, k, 5.9, True, RED)
    txt(RX + 0.34, ry2 - 0.005, RW - 0.34, 0.30, t, 6.2, False, BODY, ls=1.1)
    ry2 += 0.32

# ── 右栏：判断与边界
shead(RX, 5.42, RW, "事实与判断")
txt(RX, 5.64, RW, 0.74,
    "事实：ATLAS 由 ARINC 提出并由 IEEE 416 / 716 标准化；IEEE 1641 与 ATML 为其后续。\n"
    "判断：我们的 TPS 与之属形制同源（动作—信号—判据三要素），但无公开证据表明存在代码级移植；"
    "半导体侧语言为厂商私有，不存在可直接对标的公开标准。",
    6.3, False, RGBColor(0x3A, 0x38, 0x34), ls=1.24)

# ── 小结
hline(0.56, 6.44, 12.20, RULE, 0.010)
txt(0.56, 6.50, 0.52, 0.20, "小结", 8.0, True, RED)
txt(1.12, 6.50, 11.60, 0.44,
    "类 DSL 的 TPS 并非某家厂商的独创：其形制可追溯到 ARINC 的 ATLAS，并经 IEEE 416 / 716 标准化、由 1641 与 ATML 延续；"
    "对外表述时注明这一谱系，比笼统称为自研 DSL 更有依据。",
    7.6, False, RGBColor(0x25, 0x25, 0x25), ls=1.2)

txt(0.56, 7.09, 11.5, 0.20,
    "资料来源：IEEE 标准页（416、716、1641）、ATML 相关文献、Teradyne 与 Advantest 公开产品资料；"
    "与本文 TPS 的对应关系为归纳判断。", 6.2, False, DIM)
txt(12.30, 7.14, 0.50, 0.20, "2 / 2", 8, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
pics = oob = ov = 0
for si, s in enumerate(p.slides):
    bx = []
    for sh in s.shapes:
        if sh.shape_type == 13:
            pics += 1
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界 p{si+1}: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bx.append((l, t, w, h, sh.text_frame.text[:12]))
    for i in range(len(bx)):
        for j in range(i + 1, len(bx)):
            a, b = bx[i], bx[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠 p{si+1}: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜图片 {pics}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
