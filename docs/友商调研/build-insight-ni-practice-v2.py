# -*- coding: utf-8 -*-
"""友商调研 · NI 洞察重写 v2：图标拓扑（左）+ 三大组件（左下）+ 两段完整文字（右）"""
import os
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
IMGD = os.path.join(HERE, "ni_img")
ICON = os.path.join(IMGD, "icons")
OUT = os.path.join(HERE, "业界方案洞察-NI-洞察重写-v2.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2E, 0x2E, 0x2E)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
BLUE2 = RGBColor(0x1B, 0x5E, 0xA8)
RULE = RGBColor(0xAD, 0xA8, 0x9C)
GRID = RGBColor(0xC6, 0xC1, 0xB5)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WIRE = RGBColor(0x3C, 0x3A, 0x36)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
sl = prs.slides.add_slide(prs.slide_layouts[6])


def _font(run, size, bold, color, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.14, italic=False):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = ln
        _font(r, size, bold, color, italic)
    return tb


def hi(x, y, w, h, text, size=6.3, color=BODY, accent=BLUE2, bold=True,
       ls=1.30, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = ls
    for i, seg in enumerate(str(text).split("**")):
        if not seg:
            continue
        on = (i % 2 == 1)
        r = p.add_run()
        r.text = seg
        _font(r, size, bold if on else False, accent if on else color)
    return tb


def rect(x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(x, y, w, color=RULE, lw=0.010):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(x, y, h, color=RED, w=0.040):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def wire(pts, color=WIRE, lw=0.75, arrow=False):
    for i in range(len(pts) - 1):
        x1, y1 = pts[i]
        x2, y2 = pts[i + 1]
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(lw)
        if arrow and i == len(pts) - 2:
            ln = c.line._get_or_add_ln()
            t = etree.SubElement(ln, qn("a:tailEnd"))
            t.set("type", "triangle")
            t.set("w", "sm")
            t.set("len", "sm")


def picon(x, y, maxw, maxh, fname):
    path = fname if os.path.isabs(fname) else os.path.join(ICON if fname.endswith(".png") else IMGD, fname)
    iw, ih = PILImage.open(path).size
    ar = iw / ih
    w = maxw
    h = w / ar
    if h > maxh:
        h = maxh
        w = h * ar
    sl.shapes.add_picture(path, Inches(x + (maxw - w) / 2), Inches(y + (maxh - h) / 2),
                          width=Inches(w), height=Inches(h))


def inode(x, y, w, h, label, icon, accent=BLUE, fill=WHITE, lsize=5.4):
    rect(x, y, w, h, fill, accent, 0.8)
    vrule(x, y, h, accent, 0.032)
    ih = h - 0.055 - 0.165
    picon(x + 0.03, y + 0.045, w - 0.06, ih, icon)
    txt(x + 0.03, y + h - 0.175, w - 0.06, 0.15, label, lsize, True, INK, PP_ALIGN.CENTER)


def head(x, y, w, text, color=RED, size=9.4):
    txt(x, y, w, 0.19, text, size, True, color)
    hline(x, y + 0.215, w, GRID, 0.010)


# ══════════ 页头 ══════════
txt(0.56, 0.16, 8.20, 0.34, "业界方案洞察：NI 在 ATE 领域的优秀实践", 19, True, INK)
txt(9.00, 0.24, 3.76, 0.22, "软件拓扑 · 设备虚拟化", 8.4, False, DIM, PP_ALIGN.RIGHT)
hline(0.56, 0.54, 12.20, RULE, 0.014)

# ══════════ 洞察结论（灰底，一段完整总结）══════════
rect(0.56, 0.62, 12.20, 0.96, GRAY, None, 0)
vrule(0.56, 0.62, 0.96, RED, 0.048)
txt(0.78, 0.70, 3.00, 0.17, "洞察结论", 6.8, True, RED)
hi(0.78, 0.90, 11.86, 0.62,
   "NI 在 ATE 行业的实践可以概括为一句：**把测试能力从专用硬件搬到软件与标准接口之上**。"
   "其上以 TestStand 管理测试执行与多工位调度、以 LabVIEW 承载测量程序、以 TSM 将半导体工位与引脚概念引入执行层，"
   "其下以 NI-VISA 与 IVI 统一仪器访问、以 PXI 平台承载通道与同步；层间**单向依赖、同层可替换**，"
   "使更换仪器与新增机型不必改写测试程序。与之配套的**设备虚拟化**让仪器能被描述、被仿真、被类驱动映射，"
   "因而无实机亦可联调、跨形态只需配置。两者共同的作用是把**机型切换成本与联调周期**从硬件侧移到软件侧。",
   7.5, RGBColor(0x25, 0x25, 0x25), BLUE2, True, 1.34)

# ══════════ 中线 ══════════
CX = 6.66
rect(CX, 1.72, 0.012, 5.10, RGBColor(0x1A, 0x1A, 0x1A), None, 0)

# ══════════ 左：图标拓扑（不越中线，右界 6.50）══════════
LX = 0.56
head(LX, 1.70, 5.80, "NI 的系统拓扑（自下而上：被测件至产线系统）")

# 行 A：数据与业务
rA = [(0.56, 1.34, "MES 产线系统", "mes.png", RED, FILL2),
      (2.02, 1.24, "程序库", "lib.png", RED, WHITE),
      (3.38, 1.24, "结果库", "db.png", RED, FILL2),
      (4.74, 1.76, "分析平台", "chart.png", RED, WHITE)]
YA, HA = 1.96, 0.44
for x, w, lb, ic, ac, fl in rA:
    inode(x, YA, w, HA, lb, ic, ac, fl)

# 行 B：工业以太网交换机（通栏细条）
YB, HB = 2.50, 0.26
rect(1.10, YB, 5.10, HB, FILL2, BLUE, 0.7)
txt(1.10, YB + 0.055, 5.10, 0.16, "工业以太网交换机（千兆 · 环网 · VLAN）", 5.6, True, INK, PP_ALIGN.CENTER)

# 行 C：工控机
YC, HC = 2.86, 0.46
inode(1.10, YC, 5.10, HC, "工业控制计算机（Host）：测试执行 · 工位扩展 · 操作界面 · 驱动",
      "pc.png", BLUE, WHITE)

# 行 D：测量层
YD, HD = 3.42, 0.44
rD = [(0.56, 1.90, "测量程序（LabVIEW）", "vi.png", BLUE, WHITE),
      (2.62, 1.72, "仪器驱动（SCPI / IVI）", "driver.png", BLUE, FILL2),
      (4.50, 2.00, "开关与路径（资源仲裁）", "crossbar.png", BLUE, WHITE)]
for x, w, lb, ic, ac, fl in rD:
    inode(x, YD, w, HD, lb, ic, ac, fl)

# 行 E：仪器层
YE, HE = 3.96, 0.46
rE = [(0.56, 1.90, "模块化机箱（PXI）", "chassis_s.jpg", BLUE, WHITE),
      (2.62, 1.72, "开关与矩阵", "switch.png", BLUE, FILL2),
      (4.50, 2.00, "外挂仪器（LXI / GPIB / USB）", "instr.png", BLUE, WHITE)]
for x, w, lb, ic, ac, fl in rE:
    inode(x, YE, w, HE, lb, ic, ac, fl)

# 行 F：被测层
YF, HF = 4.50, 0.38
inode(0.56, YF, 2.90, HF, "夹具与针床", "fixture.png", RED, FILL2, 5.2)
inode(3.84, YF, 2.66, HF, "被测件工位 DUT × N", "dut.png", RED, WHITE, 5.2)

# 连线
for x, w, *_ in rA:
    wire([(x + w / 2, YA + HA), (x + w / 2, YB)], lw=0.7)
wire([(3.65, YB + HB), (3.65, YC)], lw=0.75, arrow=True)
wire([(3.65, YC + HC), (3.65, YD)], lw=0.75, arrow=True)
for x, w, *_ in rD:
    wire([(x + w / 2, YD + HD), (x + w / 2, YE)], lw=0.7)
wire([(2.01, YF - 0.04), (2.01, YE + HE)], lw=0.7)
wire([(5.17, YF - 0.04), (5.17, YE + HE)], lw=0.7)
wire([(3.46, YF + HF / 2), (3.84, YF + HF / 2)], lw=0.7)

txt(0.56, 4.92, 5.80, 0.16,
    "连线为通路示意：控制与数据（TCP/IP、PCIe、背板触发）、信号（电缆与针床接触）；机箱为产品图，其余图形为示意。",
    5.3, False, DIM)

# ══════════ 左下：三大组件（官方描述 + 软件图标）══════════
head(LX, 5.14, 5.80, "三大组件：官方功能描述（引自 ni.com）", color=BLUE)
comps = [
    ("1", "TestStand", "ni_TestStand.jpg",
     "Off-the-shelf test management software for validation & production test.",
     "现成的测试管理软件，用于验证测试与生产测试。"),
    ("2", "LabVIEW", "ni_LabVIEW.jpg",
     "A graphical programming environment engineers use to develop automated production, validation & research test systems.",
     "工程师用于开发自动化生产、验证与研究类测试系统的图形化编程环境。"),
    ("3", "TSM", "ni_TSM.jpg",
     "Extends the TestStand environment to help you develop, debug, deploy, and maintain semiconductor test systems.",
     "在 TestStand 环境之上扩展，用于开发、调试、部署与维护半导体测试系统。"),
]
cy = 5.38
for no, nm, ic, en, cn in comps:
    rect(0.56, cy, 0.50, 0.44, WHITE, GRID, 0.6)
    picon(0.58, cy + 0.02, 0.46, 0.40, ic)
    txt(1.14, cy - 0.01, 0.28, 0.16, no, 6.4, True, RGBColor(0xC9, 0xC4, 0xB8))
    txt(1.40, cy - 0.02, 1.20, 0.16, nm, 7.0, True, INK)
    txt(1.40, cy + 0.16, 4.96, 0.14, en, 5.2, False, RGBColor(0x4A, 0x48, 0x44), ls=1.05, italic=True)
    txt(1.40, cy + 0.31, 4.96, 0.16, cn, 5.6, False, BODY, ls=1.05)
    cy += 0.50

# ══════════ 右：两段完整文字 ══════════
RX, RW = 6.94, 5.82

head(RX, 1.74, RW, "NI 的软件拓扑思想", color=RED, size=11.5)
hi(RX, 2.12, RW, 1.72,
   "NI 把测试软件分成五层：最上面是产线与数据，接着是管测试流程的，再往下是写测试程序的、连仪器的，"
   "最下面是机箱与仪器本身。规则只有一条——**上层调用紧邻的下层，下层不反过来依赖上层**。"
   "这样换一台仪器，只要最下面的接口没变，上面的测试程序一行都不用改；增加一个机型，也只需换配置。"
   "测量结果以统一格式交给产线系统，测试程序不负责出报表。",
   12.0, RGBColor(0x2A, 0x2A, 0x2A), BLUE2, True, 1.70)

head(RX, 4.10, RW, "NI 的设备虚拟化思想", color=RED, size=11.5)
hi(RX, 4.48, RW, 1.85,
   "NI 把仪器当作**一组能力**来看：几个通道、多大范围、多快的采样、怎么触发，而不是一个具体型号。"
   "程序按能力写，具体是哪台仪器，由驱动去对应。于是没有实机也能先用仿真设备把流程跑通，联调不必等设备到货；"
   "要换仪器，只改驱动与配置；同一套程序在机箱模块和台式仪器上都能用。仪器的差别被收在**驱动和配置两处**，"
   "其余部分只看见能力。",
   12.0, RGBColor(0x2A, 0x2A, 0x2A), BLUE2, True, 1.70)

# 页脚
txt(0.56, 6.98, 10.9, 0.20,
    "资料来源：NI 公开产品资料与技术文档（TestStand、LabVIEW、TSM、NI-VISA、IVI、PXI）；英文引文与产品图标引自 ni.com，中文为直译；架构归纳与洞察结论为本文观点。",
    5.8, False, DIM)
txt(12.28, 7.02, 0.52, 0.20, "1 / 1", 7.4, False, DIM, PP_ALIGN.RIGHT)

prs.save(OUT)

p = Presentation(OUT)
oob = ov = 0
CXL = 6.66
for si, s in enumerate(p.slides):
    bxs = []
    for sh in s.shapes:
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if l + w > CXL + 0.02 and 1.70 < t < 6.80 and l < CXL - 0.02:
            print(f"  越中线: {sh.shape_type} ({l:.2f},{t:.2f},{w:.2f}x{h:.2f}) {sh.text_frame.text[:10] if sh.has_text_frame else ''}")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bxs.append((l, t, w, h, sh.text_frame.text[:12]))
    for i in range(len(bxs)):
        for j in range(i + 1, len(bxs)):
            a, b = bxs[i], bxs[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
