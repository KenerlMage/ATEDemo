# -*- coding: utf-8 -*-
"""业界方案洞察 · NI（拓扑版）：左栏改为节点 + 连线拓扑，术语按软件工程口径
P1 硬件拓扑：中心层 → 网络层 → 控制层 → 仪器层 → 被测层
P2 软件架构：应用层 → 执行层 → 测量层 → 平台层"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
IMGD = os.path.join(HERE, "ni_img")
OUT = os.path.join(HERE, "业界方案洞察-NI-拓扑版.pptx")

SW, SH = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x30, 0x30, 0x30)
MUT = RGBColor(0x6E, 0x6C, 0x68)
DIM = RGBColor(0x93, 0x91, 0x8C)
RED = RGBColor(0x8B, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x3F, 0x66)
RULE = RGBColor(0xBF, 0xBB, 0xB1)
FILL = RGBColor(0xF7, 0xF6, 0xF3)
FILL2 = RGBColor(0xF1, 0xF3, 0xF6)
GRAY = RGBColor(0xED, 0xED, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WIRE = RGBColor(0x7A, 0x76, 0x6E)


def crop(src, dst, ar, focus=0.5, vfocus=0.5):
    im = Image.open(os.path.join(IMGD, src)).convert("RGB")
    w, h = im.size
    tw, th = w, w / ar
    if th > h:
        th, tw = h, h * ar
    left = max(0, min(w - tw, w * focus - tw / 2))
    top = max(0, min(h - th, h * vfocus - th / 2))
    im.crop((int(left), int(top), int(left + tw), int(top + th))).save(
        os.path.join(IMGD, dst), "JPEG", quality=90, optimize=True)


crop("pxi_config.png", "rep_pxi.jpg", 3.82, 0.52, 0.5)
crop("pxi_config.png", "node_chassis.jpg", 1.55, 0.45, 0.5)
crop("labview_hero.png", "node_lv.jpg", 1.55, 0.50, 0.5)
crop("teststand_hero.png", "rep_ts.jpg", 2.00, 0.44, 0.5)
crop("labview_hero.png", "rep_lv.jpg", 2.00, 0.50, 0.5)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "微软雅黑")


def txt(sl, x, y, w, h, text, size=8, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.15):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = ln
        _font(r, size, bold, color)
    return tb


def rect(sl, x, y, w, h, fill=WHITE, line=RULE, lw=0.75):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def hline(sl, x, y, w, color=RULE, lw=0.012):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(lw))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def vrule(sl, x, y, h, color=RED, w=0.045):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def wire(sl, pts, color=WIRE, lw=1.0, arrow=True, dash=False):
    """折线：pts = [(x,y), ...]"""
    for i in range(len(pts) - 1):
        x1, y1 = pts[i]
        x2, y2 = pts[i + 1]
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(lw)
        ln = c.line._get_or_add_ln()
        if dash:
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "dash")
        if arrow and i == len(pts) - 2:
            t = etree.SubElement(ln, qn("a:tailEnd"))
            t.set("type", "triangle")
            t.set("w", "med")
            t.set("len", "med")


def node(sl, x, y, w, h, title, sub="", accent=BLUE, fill=WHITE, img=None,
         imgw=0.62, tags=""):
    rect(sl, x, y, w, h, fill, accent, 0.9)
    vrule(sl, x, y, h, accent, 0.045)
    tx = x + 0.12
    if img:
        ih = min(h - 0.14, imgw / 1.55)
        sl.shapes.add_picture(os.path.join(IMGD, img), Inches(x + 0.13),
                              Inches(y + (h - ih) / 2), width=Inches(imgw), height=Inches(ih))
        rect(sl, x + 0.13, y + (h - ih) / 2, imgw, ih, None, RULE, 0.6)
        tx = x + 0.13 + imgw + 0.10
    tw = x + w - tx - 0.10 - (0.54 if tags else 0.0)
    txt(sl, tx, y + 0.09, tw, 0.19, title, 7.3, True, INK)
    if sub:
        txt(sl, tx, y + 0.29, tw, h - 0.33, sub, 6.1, False, MUT, ls=1.12)
    if tags:
        txt(sl, x + w - 0.60, y + 0.10, 0.46, 0.16, tags, 5.8, False, DIM, PP_ALIGN.RIGHT)


def ltag(sl, x, y, text, w=1.30, size=5.8):
    txt(sl, x, y, w, 0.14, text, size, False, RGBColor(0x86, 0x82, 0x7A))


def head(sl, sub_lead, sub_rest, page):
    txt(sl, 0.56, 0.24, 8.0, 0.38, "业界方案洞察", 21, True, INK)
    txt(sl, 8.90, 0.36, 3.90, 0.22, sub_lead, 8.6, False, DIM, PP_ALIGN.RIGHT)
    rect(sl, 0.56, 0.70, 12.20, 0.64, GRAY, None, 0)
    vrule(sl, 0.56, 0.70, 0.64, RED, 0.05)
    txt(sl, 0.74, 0.80, 11.85, 0.48, sub_rest, 7.6, False, RGBColor(0x2E, 0x2E, 0x2E), ls=1.25)
    txt(sl, 12.30, 7.14, 0.50, 0.20, page, 8, False, DIM, PP_ALIGN.RIGHT)


def rhead(sl, x, y, w, text):
    txt(sl, x, y, w, 0.18, text, 8.2, True, RED)
    hline(sl, x, y + 0.23, w, RULE)


def scene(sl, x, y, w, rows, pitch=0.30):
    for i, (a, b) in enumerate(rows):
        rect(sl, x, y + i * pitch, w, pitch, WHITE if i % 2 else FILL, RULE, 0.5)
        txt(sl, x + 0.10, y + i * pitch + 0.07, w - 0.92, 0.18, a, 7.0, False, BODY)
        txt(sl, x + w - 0.72, y + i * pitch + 0.07, 0.62, 0.18, b, 7.0, True,
            RED if b == "受限" else INK, PP_ALIGN.RIGHT)


def summary(sl, text):
    rect(sl, 0.56, 6.34, 12.20, 0.58, WHITE, RULE, 0.75)
    vrule(sl, 0.56, 6.34, 0.58, RED, 0.05)
    txt(sl, 0.70, 6.41, 0.52, 0.20, "小结", 8.0, True, RED)
    txt(sl, 1.22, 6.42, 11.40, 0.44, text, 7.8, False, RGBColor(0x25, 0x25, 0x25), ls=1.2)


def foot(sl, s):
    hline(sl, 0.56, 7.02, 12.20, RULE, 0.010)
    txt(sl, 0.56, 7.09, 11.5, 0.20, s, 6.6, False, DIM)


def pic(sl, path, x, y, w, h, cap):
    sl.shapes.add_picture(os.path.join(IMGD, path), Inches(x), Inches(y),
                          width=Inches(w), height=Inches(h))
    rect(sl, x, y, w, h, None, RULE, 0.75)
    txt(sl, x, y + h + 0.03, w, 0.18, cap, 6.4, False, DIM)


ABOUT_HW = ("美国国家仪器（National Instruments，简称 NI）成立于 1976 年，总部位于美国得克萨斯州奥斯汀，"
            "长期提供模块化仪器、测试执行软件与系统集成工具；2023 年并入 Emerson，作为其测试与测量业务运营。"
            "其 PXI 模块化平台与 TestStand 测试执行软件在电子、汽车与半导体测试领域应用广泛。"
            "本页基于其公开资料，说明该方案的硬件拓扑与适用场景。")
ABOUT_SW = ("NI 的软件体系以 LabVIEW（1986 年首次发布）为基础，向上延伸出 TestStand 测试执行软件，"
            "以及面向半导体测试的 TestStand Semiconductor Module。其公开资料将测量、测试执行与数据管理划分为"
            "相对独立的产品线，各自独立发布与升级。本页说明该方案的软件分层与适用场景。")

RX, RW = 8.02, 4.74

# ══════════════════════════ P1 硬件拓扑
s1 = prs.slides.add_slide(BLANK)
head(s1, "第一篇　硬件拓扑", ABOUT_HW, "1 / 2")

# 中心层
node(s1, 0.56, 1.52, 1.78, 0.56, "MES 产线系统", "任务令 · 结果回传", RED, FILL2)
node(s1, 2.50, 1.52, 1.62, 0.56, "程序库", "序列与配置版本", RED, WHITE)
node(s1, 4.28, 1.52, 1.62, 0.56, "结果库", "原始数据 · 追溯", RED, FILL2)
node(s1, 6.06, 1.52, 1.80, 0.56, "分析平台", "报表 · 良率趋势", RED, WHITE)

# 中心层汇入交换机
wire(s1, [(1.45, 2.08), (1.45, 2.24)], arrow=False)
wire(s1, [(3.31, 2.08), (3.31, 2.24)], arrow=False)
wire(s1, [(5.09, 2.08), (5.09, 2.24)], arrow=False)
wire(s1, [(6.96, 2.08), (6.96, 2.24)], arrow=False)
wire(s1, [(1.45, 2.24), (6.96, 2.24)], arrow=False)

node(s1, 2.62, 2.32, 3.60, 0.54, "工业以太网交换机", "千兆 · 环网冗余 · VLAN 隔离", BLUE, FILL2, tags="L3")
wire(s1, [(4.42, 2.24), (4.42, 2.32)], arrow=False)
wire(s1, [(4.42, 2.86), (4.42, 3.02)])
ltag(s1, 4.50, 2.88, "TCP/IP 千兆", 1.2)

# 控制层
node(s1, 1.90, 3.04, 5.10, 0.62, "工业控制计算机（Host）",
     "测试执行软件 · 测量程序 · 操作界面 · 设备驱动", BLUE, WHITE, tags="L2")
wire(s1, [(3.30, 3.66), (3.30, 3.84)], arrow=False)
wire(s1, [(1.86, 3.84), (6.96, 3.84)], arrow=False)
wire(s1, [(1.86, 3.84), (1.86, 4.10)])
wire(s1, [(4.27, 3.84), (4.27, 4.10)])
wire(s1, [(6.59, 3.84), (6.59, 4.10)])
ltag(s1, 5.10, 3.66, "PCIe / 以太网　背板触发", 2.0, 5.8)

# 仪器层
node(s1, 0.56, 4.12, 2.60, 0.86, "模块化机箱", "数字 / 模拟 / 射频模块", BLUE, WHITE,
     img="node_chassis.jpg", imgw=0.68, tags="L1")
node(s1, 3.42, 4.12, 1.70, 0.86, "开关与矩阵", "通道扩展\n路径互斥", BLUE, FILL2, tags="L1")
node(s1, 5.32, 4.12, 2.54, 0.86, "外挂仪器", "LXI / GPIB / USB\n统一会话访问", BLUE, WHITE, tags="L1")

# 仪器层到被测层：总线 + 分支
wire(s1, [(1.86, 4.98), (1.86, 5.18)], arrow=False)
wire(s1, [(4.27, 4.98), (4.27, 5.18)], arrow=False)
wire(s1, [(6.59, 4.98), (6.59, 5.18)], arrow=False)
wire(s1, [(1.86, 5.18), (6.59, 5.18)], arrow=False)
wire(s1, [(2.60, 5.18), (2.60, 5.34)])
wire(s1, [(6.36, 5.18), (6.36, 5.34)])
ltag(s1, 3.42, 5.20, "信号电缆 / 背板触发总线", 1.9, 5.8)

# 被测层
node(s1, 0.56, 5.36, 4.10, 0.58, "夹具与针床", "机械定位 · 信号引出 · 品种更换整体替换", RED, FILL2)
node(s1, 4.86, 5.36, 3.00, 0.58, "被测件工位 DUT × N", "并行执行 · 手动或自动上下料", RED, WHITE)
wire(s1, [(4.66, 5.65), (4.86, 5.65)])
ltag(s1, 4.58, 5.70, "接触", 0.34, 5.6)

pic(s1, "rep_pxi.jpg", RX, 1.50, RW, 1.24, "图 1　模块化机箱与仪器配置界面（ni.com 产品资料）")

rhead(s1, RX, 2.96, RW, "拓扑说明")
notes = [
    ("（一）", "单站点纵向结构：中心层承担程序下发与结果回收，控制层承担全部实时测试动作。"),
    ("（二）", "多工位由机箱内并行执行实现，不通过增加控制计算机数量扩展工位；"
              "程序与结果走工业以太网，仪器控制信号走背板触发与专用电缆，两条通路相互独立。"),
]
ny = 3.24
for k, t in notes:
    txt(s1, RX, ny, 0.42, 0.18, k, 7.2, True, DIM)
    txt(s1, RX + 0.42, ny, RW - 0.42, 0.56, t, 7.2, False, BODY, ls=1.2)
    ny += 0.62

rhead(s1, RX, 4.50, RW, "适用场景")
scene(s1, RX, 4.76, RW, [
    ("多品种、中小批量、引脚数中等以上", "适用"),
    ("单一机箱内 4 至 16 工位并行", "适用"),
    ("数字、模拟与射频混合测量", "适用"),
    ("成本敏感的单品种大批量", "受限"),
    ("需与既有非 NI 仪器深度耦合", "受限"),
])

summary(s1, "该方案以单一控制节点承载多工位并行，扩展工位不依赖增加计算机数量；程序与结果、仪器控制分走两条通路。"
            "适用于中等规模、品种较多的量产线；适用性判断以功能覆盖与工程成本为依据，不含价格因素。")
foot(s1, "资料来源：NI 公开产品资料与公司信息；节点划分与适用性判断为本文归纳。")

# ══════════════════════════ P2 软件架构
s2 = prs.slides.add_slide(BLANK)
head(s2, "第二篇　软件架构", ABOUT_SW, "2 / 2")

# 应用层
node(s2, 0.56, 1.52, 2.30, 0.58, "操作界面（HMI）", "启动 · 监视 · 异常处置", BLUE, FILL2, tags="应用")
node(s2, 3.00, 1.52, 2.26, 0.58, "程序管理", "序列与版本下发", BLUE, WHITE, tags="应用")
node(s2, 5.36, 1.52, 2.50, 0.58, "结果查询与报表", "合格率 · 追溯查询", BLUE, FILL2, tags="应用")

wire(s2, [(1.71, 2.10), (1.71, 2.34), (3.10, 2.34), (3.10, 2.56)])
wire(s2, [(4.13, 2.10), (4.13, 2.56)])
wire(s2, [(6.61, 2.10), (6.61, 2.34), (5.60, 2.34), (5.60, 2.56)])
ltag(s2, 2.10, 2.16, "调用接口", 1.10, 5.8)

# 执行层
node(s2, 0.56, 2.58, 3.24, 0.66, "测试执行引擎", "流程组织 · 步骤调度 · 超时与异常", RED, WHITE, tags="执行")
node(s2, 3.98, 2.58, 1.86, 0.66, "工位管理", "多工位并行\n合格性判定与分档", RED, FILL2, tags="执行")
node(s2, 6.00, 2.58, 1.86, 0.66, "结果归档", "统一格式\n批次与序列号", RED, WHITE, tags="执行")

wire(s2, [(3.80, 2.91), (3.98, 2.91)])
wire(s2, [(5.84, 2.91), (6.00, 2.91)], dash=False)

# 执行层到测量层
wire(s2, [(2.18, 3.24), (2.18, 3.48), (2.10, 3.48), (2.10, 3.74)])
wire(s2, [(4.91, 3.24), (4.91, 3.74)])
wire(s2, [(6.93, 3.24), (6.93, 3.74)])
ltag(s2, 5.00, 3.34, "调用测量接口", 1.30, 5.8)

# 测量层
node(s2, 0.56, 3.76, 3.02, 0.78, "测量程序", "激励施加 · 信号采集\n数据换算与限值比较", BLUE, WHITE,
     img="node_lv.jpg", imgw=0.64, tags="测量")
node(s2, 3.76, 3.76, 2.00, 0.78, "仪器驱动", "会话管理\nSCPI · IVI · 寄存器", BLUE, FILL2, tags="测量")
node(s2, 5.92, 3.76, 1.94, 0.78, "开关与路径", "通路互斥\n共享资源仲裁", BLUE, WHITE, tags="测量")

wire(s2, [(4.62, 4.54), (4.62, 4.82), (4.10, 4.82), (4.10, 5.04)])
wire(s2, [(6.89, 4.54), (6.89, 5.04)])
ltag(s2, 5.10, 4.60, "资源请求", 0.90, 5.8)

# 平台层
node(s2, 0.56, 5.06, 3.24, 0.66, "资源配置与自检", "通道映射 · 时钟与触发 · 固件", BLUE, FILL2, tags="平台")
node(s2, 3.98, 5.06, 1.86, 0.66, "机箱与模块管理", "资源发现\n设备拓扑", BLUE, WHITE, tags="平台")
node(s2, 6.00, 5.06, 1.86, 0.66, "设备维护", "校准记录\n备件与生命周期", BLUE, FILL2, tags="平台")

wire(s2, [(3.80, 5.39), (3.98, 5.39)], arrow=False)
wire(s2, [(5.84, 5.39), (6.00, 5.39)], arrow=False)
ltag(s2, 1.20, 5.74, "层间以配置与调用约定衔接，箭头表示调用方向", 3.80, 5.8)

pic(s2, "rep_ts.jpg", RX, 1.50, 2.32, 1.16, "图 1　测试执行软件")
pic(s2, "rep_lv.jpg", RX + 2.42, 1.50, 2.32, 1.16, "图 2　测量开发环境")

rhead(s2, RX, 2.88, RW, "架构特征")
feats = [
    ("（一）", "层次职责单一：应用层不涉及测量，测量层不涉及流程，任一层调整不影响相邻层。"),
    ("（二）", "流程与测量分离、结果出口统一：换被测对象主要改测量层与流程配置，换仪器主要改测量层；"
              "各层结果按统一格式归档，供中心侧汇总与追溯。"),
]
ny = 3.16
for k, t in feats:
    txt(s2, RX, ny, 0.42, 0.18, k, 7.2, True, DIM)
    txt(s2, RX + 0.42, ny, RW - 0.42, 0.58, t, 7.2, False, BODY, ls=1.2)
    ny += 0.64

rhead(s2, RX, 4.50, RW, "适用场景")
scene(s2, RX, 4.76, RW, [
    ("流程稳定、需长期维护的量产线", "适用"),
    ("测试项由工艺定义、软件仅作组织", "适用"),
    ("多品种小批量、换型频繁", "适用"),
    ("频繁变更的研发验证阶段", "受限"),
    ("仅需轻量单工位执行的小型工位", "受限"),
])

summary(s2, "软件按应用、执行、测量、平台四层划分，流程与测量分离使换型与换仪器的影响范围可控；"
            "该结构适合流程稳定、需长期维护的量产场景，研发验证阶段因配置与授权调整周期较长而相对受限。")
foot(s2, "资料来源：NI 公开产品资料与公司信息；四层划分与适用性判断为本文归纳，非厂商官方分类。")

prs.save(OUT)

p = Presentation(OUT)
pics = oob = ov = 0
for si, sl in enumerate(p.slides):
    bx = []
    for sh in sl.shapes:
        if sh.shape_type == 13:
            pics += 1
        try:
            l, t = sh.left.inches, sh.top.inches
            w, h = sh.width.inches, sh.height.inches
        except Exception:
            continue
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            oob += 1
            print(f"  越界 p{si+1}: ({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            bx.append((l, t, w, h, sh.text_frame.text[:12]))
    for i in range(len(bx)):
        for j in range(i + 1, len(bx)):
            a, b = bx[i], bx[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.08 and oy > 0.08:
                ov += 1
                print(f"  重叠 p{si+1}: 「{a[4]}」×「{b[4]}」{ox:.2f}x{oy:.2f}")
print(f"页数 {len(p.slides)}｜图片 {pics}｜越界 {oob}｜文本重叠 {ov}｜{os.path.getsize(OUT)} B")
